#!/usr/bin/env python3
"""Every Logical Framework field at the EXACT size the platform allows.

    python3 build_form_fields.py

His screenshots gave me the real field limits, and three of them are far
smaller than anything I had assumed:

    Specific Objective        300     free text, not a dropdown label
    Expected result           300     MINE WAS 381 - it did not fit
    Output Title              100     two of my eight titles were over
    Output Description        500     ALL EIGHT of mine were over, some double
    Measurement unit           50     two of mine were over
    Work package                      REQUIRED, multi-select
    Semester of delivery              dropdown
    Target value                      number

THE ONE THAT MATTERS MOST: THE EXPECTED RESULT I SENT WAS TOO LONG

I wrote 381 characters for a 300-character box and told him it was deliberately
short. It was deliberately short for the wrong limit. If he pasted it, the
platform will have cut it at 300 - mid-sentence, in the single statement every
indicator attaches to.

The row in his screenshot reads "Receiving Mediterranean territories apply a
transferred and ..." which is the table truncating for display, so I cannot
tell from the picture whether the stored text is whole. He has to open Edit and
look. The version below is 300 or under and says the same three things.

That is twice now that a field has been smaller than I assumed - 3.1.4 turned
out to be two boxes, and these are a sixth the size of their neighbours. The
lesson I keep relearning: ASK FOR THE COUNTER, do not infer the limit from the
field's importance. A 300-character box for the project's expected result and a
2000-character box for a methodology paragraph is not a ranking of importance,
it is just how the form was built.

WHAT THE OUTPUTS TAB ACTUALLY WANTS, WHICH I ALSO HAD WRONG

My outputs document gave each output a title, target, unit, semester and
budget. The form also wants a DESCRIPTION (500) and a WORK PACKAGE selection,
and the work package is required. The budget is not on this screen at all - it
is entered later as a percentage in the budget section.
"""
import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from _shared import here

BASE = str(here(__file__))
OUT = f"{BASE}/AMMOS_form_fields_exact_2026-10-05.docx"

NAVY = RGBColor(0x1F, 0x38, 0x5E)
GREY = RGBColor(0x5A, 0x6B, 0x7D)
RED = RGBColor(0xB4, 0x55, 0x3C)

# Measured off his screenshots, not assumed.
LIM_SO = 300
LIM_ER = 300
LIM_TITLE = 100
LIM_DESC = 500
LIM_UNIT = 50

ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV"}

RANGE = re.compile(r"(?<=\d)\s*[–—]\s*(?=[A-Za-z]?\d)")


def clean(t):
    t = str(t)
    t = RANGE.sub("\x00", t)
    t = re.sub(r"(?<!-)\s*--\s*(?!-)", ", ", t)
    t = re.sub(r"\s*[–—]\s*", ", ", t)
    t = t.replace("\x00", "-")
    return re.sub(r",\s*,", ",", t)


SPECIFIC_OBJECTIVE = (
    "Transfer and adapt a validated Mediterranean shoreline sediment "
    "management method to three territories and three shore types, "
    "demonstrate it in the field, and embed it in adopted local management "
    "plans, so that visitor pressure on sediment shores is managed as a "
    "climate adaptation measure."
)

EXPECTED_RESULT = (
    "Receiving Mediterranean territories apply an adapted shoreline sediment "
    "management method, sustain it through formally adopted costed local "
    "plans with standing multi-stakeholder governance, and make it replicable "
    "through a transferability guide and an open monitoring platform in the "
    "MMM database."
)

# (code, WP, title<=100, description<=500, target, unit<=50, semester)
FORM_OUTPUTS = [
    ("Output3.1", "WP3",
     "Modular transferability decision: 7 method sheets against 3 shore types",
     "The source method is seven field method sheets, and the receiving sites "
     "are a marine dune system, a marine coastal site and an artificial "
     "freshwater pond. Transfer is therefore 21 separate decisions, not one. "
     "Each is recorded as APPLIED, SUBSTITUTED or NOT APPLICABLE with its "
     "reason. The document carries an explicit uptake section stating what an "
     "adopting organisation must do. Establishing where a method stops working is what "
     "makes it usable by a territory that was never in the partnership.",
     21, "decisions recorded (7 sheets x 3 sites)", 2),

    ("Output3.2", "WP3",
     "Territorial, biodiversity and tourism baseline for the three "
     "demonstration sites",
     "One comparable baseline per site, produced WITH the adapted method so "
     "that the method is tested by being used. Topography, shore state, "
     "utilities and access, sea level rise scenarios "
     "from 0.2 m to 1 m, biodiversity state and ecological status. The "
     "tourism layer is part of it: visitor numbers and seasonality, a visitor "
     "survey on willingness to visit and to pay, and its economic "
     "contribution. Without that layer a baseline measures a habitat; with "
     "it, the asset the authority is deciding about.",
     3, "site baselines (one per demonstration site)", 2),

    ("Output4.1", "WP4",
     "Three demonstrated shore management installations, one per site",
     "The method applied rather than described. Visitor channelling and shore "
     "protection designed from each site's own baseline: boundary marking, "
     "signs showing what is protected and why, demountable "
     "modular boardwalks over the vulnerable surface, and demarcated zones "
     "where trampling is degrading vegetation or banks. Demountable keeps it "
     "reversible beside a protected habitat and the screening proportionate. "
     "Each is then observed through one full visitor season. No underwater "
     "cleaning at any site.",
     3, "sites with the method demonstrated in the field", 4),

    ("Output4.2", "WP4",
     "Shared cross-typology monitoring platform and common protocol",
     "One platform, one protocol, three shore types. The COMMON beach litter "
     "protocol is adopted as published so data stays comparable with the "
     "European Environment Agency dataset; the adapted sheets supply the "
     "rest. Site staff enter observations in their own language. The "
     "comparability layer is the new element: it lets a freshwater pond and a "
     "marine dune be read side by side, and carries an uptake section saying "
     "what a new territory must do to join. Open licensing and data access "
     "are fixed here.",
     1, "web-GIS platform serving three territories", 3),

    ("Output5.1", "WP5",
     "Shore Custodians adapted to a lake shore and a coastal dune, run with "
     "schools",
     "The AMMIRARE Beach Custodians instrument transferred to shore types it "
     "was not written for - the same adaptation problem as the method "
     "sheets. Teaching material in Greek, Turkish and "
     "Arabic, with field exercises reworked to what each site actually has. "
     "Delivered with the schools nearest each site so the groups survive the "
     "project by being local to it. Each group runs at least one monitoring "
     "round and one awareness action, and its own observations enter the "
     "shared platform.",
     9, "school custodian groups (3 per country)", 3),

    ("Output5.2", "WP5",
     "Site-manager training and visitor awareness package",
     "The people who will still be there in year three. Training built on the "
     "COMMON Capacity Building Training Design Plan, delivered in each country "
     "in the local language with a practical module at the demonstration "
     "site. Aimed at municipal "
     "and authority staff holding a mandate over the site, so that the "
     "management plans are written by people who have used the method. "
     "Visitor awareness adapts the BEach CLEAN decalogue, which already "
     "exists in Arabic and is extended rather than translated afresh.",
     60, "site managers and local staff trained", 4),

    ("Output6.1", "WP6",
     "Three costed site management plans taken through consultation to "
     "adoption",
     "A plan that is filed changes nothing. Each is costed, so the authority "
     "knows what sustaining it requires, and consulted to a formal adoption "
     "decision by the body with the mandate. A Shore Stewardship Committee is "
     "constituted alongside each plan - the responsible authority, a tourism "
     "operator, the research partner, a civil society or youth voice - so "
     "every plan has an owner after the project. For Lemnos the geopark route "
     "carries maintenance through an existing organisation.",
     3, "management plans formally adopted", 4),

    ("Output6.2", "WP6",
     "Joint transferability guide and policy recommendations, English and "
     "Arabic",
     "What a fourth territory needs in order to do this without us: the "
     "adapted method, the 21 transferability decisions and their reasons, the "
     "comparability layer, and the costs. Its closing section states the "
     "actions a new territory must take to adopt or upscale the method, which "
     "makes it a countable solution, not a publication. Includes the "
     "Acipayam canyon and pond baseline as a worked example of a site type "
     "the source method never addressed. "
     "Lodged in the MMM common database under open licence.",
     1, "transferability guide lodged in MMM database", 4),
]


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Cm(1.4)
        n = self.d.styles["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(10)

    def h(self, t, size=13, colour=NAVY, space=12):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(space)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(clean(t))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), colour

    def p(self, t, size=10, colour=None, italic=False, bold=False):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run(clean(t))
        r.font.size, r.italic, r.bold = Pt(size), italic, bold
        if colour:
            r.font.color.rgb = colour

    def table(self, head, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, x in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(clean(x))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8), NAVY
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(clean(v))
                r.font.size = Pt(8)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t


def check():
    """Every field against the limit measured off his screen. No exceptions."""
    fails = []

    def lim(name, text, limit):
        n = len(clean(text))
        if n > limit:
            fails.append(f"{name} is {n} chars, limit {limit} - cut "
                         f"{n - limit}")
        return n

    n_so = lim("Specific Objective", SPECIFIC_OBJECTIVE, LIM_SO)
    n_er = lim("Expected result", EXPECTED_RESULT, LIM_ER)

    seen = set()
    for code, wp, title, desc, target, unit, semester in FORM_OUTPUTS:
        lim(f"{code} title", title, LIM_TITLE)
        lim(f"{code} description", desc, LIM_DESC)
        lim(f"{code} unit", unit, LIM_UNIT)
        if code in seen:
            fails.append(f"{code} appears twice")
        seen.add(code)
        if not (1 <= semester <= 4):
            fails.append(f"{code} semester {semester} outside I-IV")
        if not isinstance(target, int) or target <= 0:
            fails.append(f"{code} target {target!r} is not a positive whole "
                         f"number - the field is numeric")
        if wp not in ("WP3", "WP4", "WP5", "WP6"):
            fails.append(f"{code} is on {wp}; only technical work packages "
                         f"carry outputs")
        if not code.startswith(f"Output{wp[2]}."):
            fails.append(f"{code} does not match its work package {wp}")

    if len(FORM_OUTPUTS) != 8:
        fails.append(f"{len(FORM_OUTPUTS)} outputs, expected 8")

    # The single expected result has to carry all three ideas, because the
    # Guidelines allow only one and these are what the project promises.
    low = clean(EXPECTED_RESULT).lower()
    for word in ("adapt", "adopt", "replica"):
        if word not in low:
            fails.append(f"expected result does not cover {word!r}")

    blob = " ".join(str(x) for row in FORM_OUTPUTS for x in row) \
        + SPECIFIC_OBJECTIVE + EXPECTED_RESULT
    for forbidden in (r"\bRAS\b", r"Regione Autonoma", r"Hakanson", r"xxxx"):
        if re.search(forbidden, blob, re.I):
            fails.append(f"text contains {forbidden!r}")

    assert not fails, "NOT SHIPPING:\n  - " + "\n  - ".join(fails)
    return n_so, n_er


def main():
    n_so, n_er = check()

    D = Doc()
    D.h("AMMOS — every Logical Framework field at the size the platform "
        "actually allows", size=15)
    D.p(f"Generated {dt.date.today():%d/%m/%Y} from the limits in your "
        f"screenshots. Every field below is checked against its own counter "
        f"before this document is produced.", size=9.5, colour=GREY,
        italic=True)

    D.h("First — please check the Expected result you already added", size=13)
    D.p("The box is 300 characters. What I sent you yesterday was 381. I "
        "told you it was deliberately short, and it was — but short for a "
        "2000-character box, which is what I assumed. If you pasted it, the "
        "platform will have cut it at 300, mid-sentence, in the one statement "
        "every indicator attaches to.", size=10.5, bold=True)
    D.p("Your screenshot shows the row reading \"Receiving Mediterranean "
        "territories apply a transferred and ...\" but that is the table "
        "truncating for display, so I cannot tell from the picture whether "
        "the stored text is whole. Click Edit on that row and look. Replace "
        "it with the version below either way.", size=10.5)

    D.h("Specific Objective — it is a free text box, not just the code",
        size=13)
    D.p(f"{n_so} of {LIM_SO} characters. Your row currently shows \"2.2 "
        f"(RSO2.4)\" in that column, which is the programme's label rather "
        f"than a description. The field says \"Describe specific objective\" "
        f"and gives you 300 characters to do it in.", size=9,
        colour=GREY, italic=True)
    D.p(SPECIFIC_OBJECTIVE)

    D.h("Expected result — one only", size=13)
    D.p(f"{n_er} of {LIM_ER} characters. Carries transfer, adoption and "
        f"replication in one sentence, because we only get one.", size=9,
        colour=GREY, italic=True)
    D.p(EXPECTED_RESULT)

    D.h("The eight outputs, field by field", size=14)
    D.p("Title, Description, Measurement unit — all within their counters. "
        "Work package is required and is a multi-select. The budget is NOT on "
        "this screen; it is entered later as a percentage per output in the "
        "budget section, where the percentages within each work package must "
        "total exactly 100.", size=9.5, colour=GREY)
    D.table(["Code", "WP", "Semester", "Target", "Measurement unit"],
            [[c, wp, ROMAN[s], str(t), u]
             for c, wp, _, _, t, u, s in FORM_OUTPUTS],
            widths=[2.2, 1.4, 1.8, 1.6, 7.0])

    for code, wp, title, desc, target, unit, semester in FORM_OUTPUTS:
        D.h(f"{code}  —  {wp}", size=11.5, space=10)
        D.table(["Field", "Value", "Count"],
                [["Title", title, f"{len(clean(title))}/{LIM_TITLE}"],
                 ["Description", desc, f"{len(clean(desc))}/{LIM_DESC}"],
                 ["Work package", wp, "required"],
                 ["Semester of delivery", ROMAN[semester], ""],
                 ["Target value", str(target), "numeric"],
                 ["Measurement unit", unit, f"{len(clean(unit))}/{LIM_UNIT}"]],
                widths=[3.0, 12.2, 1.8])

    D.h("Then the Indicators tab", size=13)
    D.p("The codes are generated from the specific objective, so you pick "
        "rather than type. Associate and set the target:", size=10)
    D.table(["Indicator", "Target", "What carries it"],
            [["RCO84 Pilot actions developed jointly and implemented", "3",
              "Output4.1, the three demonstration installations"],
             ["RCO87 Organisations cooperating across borders", "7",
              "6 partners + 1 associated organisation - the definition counts "
              "associates, so the North Aegean statement adds a unit here"],
             ["RCO81 Participations in joint actions across borders", "60",
              "Output5.1 and Output5.2 plus the workshops. This is where the "
              "60 trained staff go - there is NO training indicator under our "
              "objective"],
             ["RCO116 Jointly developed solutions", "3",
              "Output3.1, Output4.2, Output6.2"],
             ["RCR104 Solutions taken up or up-scaled", "3",
              "the adapted METHOD taken up, evidenced by the three adopted "
              "plans - do not describe the plans themselves as the solution, "
              "the definition excludes administrative ones"],
             ["RCR84 Organisations cooperating after completion", "7",
              "all partners and the associate, under the cooperation "
              "agreement signed in WP6"]],
            widths=[6.0, 1.4, 9.6])

    D.d.save(OUT)
    print(f"wrote {OUT}")
    print(f"  Specific Objective  {n_so}/{LIM_SO}")
    print(f"  Expected result     {n_er}/{LIM_ER}")
    print(f"  {len(FORM_OUTPUTS)} outputs, all fields within limits:")
    for c, wp, t, d, tv, u, s in FORM_OUTPUTS:
        print(f"    {c} {wp}  title {len(clean(t)):>3}/{LIM_TITLE}  "
              f"desc {len(clean(d)):>3}/{LIM_DESC}  "
              f"unit {len(clean(u)):>2}/{LIM_UNIT}  sem {ROMAN[s]}  "
              f"target {tv}")


if __name__ == "__main__":
    main()
