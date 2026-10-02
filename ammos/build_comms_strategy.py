#!/usr/bin/env python3
"""AMMOS communication strategy, for Application Form section 3.4.4 and WP2.

WHY THIS IS A SEPARATE DOCUMENT AND NOT A GENERATED SECTION

Award criterion 4.4 is scored by the external evaluators at Step 1 and COUNTS
DOUBLE: 8 points of 88, the largest single item on the grid, against a
threshold of 75 where only 13 points may be lost in total.

Every applicant writes the same paragraph here - a website, social media, a
final conference, a leaflet. That paragraph scores 2 out of 4, and doubled it
costs 4 points, nearly a third of the whole allowance.

So this is drafted rather than assembled, and the thing that makes it specific
is a single idea: AMMOS IS A CAPITALISATION PROJECT, SO ITS COMMUNICATION
OBJECT IS NOT "WE DID SOMETHING", IT IS "HERE IS A METHOD YOU CAN USE". That
changes the audience, the channel and the measure of success. Dissemination
aims at awareness; transfer aims at somebody else doing it.

WHAT IS DELIBERATELY LEFT BLANK

Every reach figure is marked [FROM PARTNER]. Criterion 1.3 and 2.3 both use
the word QUANTIFIED and I will not invent an audience size to fill a gap -
twice already on this proposal an illustrative figure of mine has been
forwarded and negotiated against. The partners have the numbers; the structure
is here to receive them.
"""
import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from _shared import here

BASE = str(here(__file__))
OUT = f"{BASE}/AMMOS_communication_strategy_2026-10-02.docx"

NAVY = RGBColor(0x1F, 0x38, 0x5E)
GREY = RGBColor(0x5A, 0x6B, 0x7D)
RED = RGBColor(0xB4, 0x55, 0x3C)

# His standing rule: no dash as punctuation, ranges preserved.
RANGE = re.compile(r"(?<=\d)\s*[–—]\s*(?=[A-Za-z]?\d)")


def clean(t):
    t = str(t)
    t = RANGE.sub("\x00", t)
    t = re.sub(r"(?<!-)\s*--\s*(?!-)", ", ", t)
    t = re.sub(r"\s*[–—]\s*", ", ", t)
    t = t.replace("\x00", "-")
    return re.sub(r",\s*,", ",", t)


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Cm(2.0)
        n = self.d.styles["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(10.5)

    def h(self, t, size=13, colour=NAVY, space=12):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(space)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(clean(t))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), colour

    def p(self, t, size=10.5, colour=None, italic=False, bold=False):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(5)
        r = par.add_run(clean(t))
        r.font.size, r.italic, r.bold = Pt(size), italic, bold
        if colour:
            r.font.color.rgb = colour

    def bullet(self, t, size=10.5):
        par = self.d.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(3)
        par.add_run(clean(t)).font.size = Pt(size)

    def table(self, head, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, x in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(clean(x))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), NAVY
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(clean(v))
                r.font.size = Pt(8.5)
                if "[FROM PARTNER]" in str(v):
                    r.font.color.rgb = RED
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t


TBD = "[FROM PARTNER]"


def main():
    D = Doc()
    D.h("AMMOS, communication strategy", size=16)
    D.p("For Application Form section 3.4.4 and work package 2. "
        f"Draft of {dt.date.today():%d/%m/%Y}.", size=9.5, colour=GREY,
        italic=True)

    # --------------------------------------------------- the governing idea
    D.h("1. The objective, and why it is not dissemination")
    D.p("AMMOS does not communicate a discovery. It communicates a METHOD that "
        "already exists and has already been proven, to people who do not yet "
        "use it. The output we capitalise, MMM_IF02, is a set of seven "
        "technical sheets written for beach managers in the western "
        "Mediterranean. Nothing about them is secret and nothing about them "
        "is new. They have simply never travelled east or south.")
    D.p("So the measure of success is not how many people heard of AMMOS. It "
        "is how many site managers outside this partnership are applying the "
        "adapted method at the end of it, and whether the method is findable "
        "and usable by the next person who needs it without asking us.",
        bold=True)
    D.p("Three objectives follow from that, and every activity in work package "
        "2 serves one of them.")
    D.bullet("TRANSFER. Put the adapted method in the hands of the people who "
             "manage sediment shores in the three territories, in a language "
             "and a format they can use on a beach.")
    D.bullet("ADOPTION. Give the public authorities with territorial "
             "competence a reason and a route to adopt it, which is what "
             "turns a pilot into a management practice.")
    D.bullet("REPLICATION. Leave the method where the NEXT project will find "
             "it, which for this programme means the MMM database itself.")

    D.h("The point an evaluator will look for, stated plainly", size=11,
        space=10)
    D.p("AMMOS capitalises outputs from the MMM database and returns new ones "
        "to it. The adapted method, the comparability protocol across three "
        "shore typologies and the lacustrine adaptation are all submitted to "
        "the same database that supplied MMM_IF01 and MMM_IF02. A "
        "capitalisation project that does not feed the capitalisation "
        "mechanism has not understood the call.", bold=True)

    # ------------------------------------------------------- audiences
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("2. Audiences, named")
    D.p("The criterion asks about public authorities, tourism operators, local "
        "communities, media and policy makers. Below they are named, because a "
        "strategy addressed to categories is addressed to nobody.",
        size=9.5, colour=GREY)
    D.table(
        ["Audience", "Who, specifically", "What they must be able to do after",
         "Reach"],
        [["Site managers, the primary audience",
          "Acıpayam Municipality at Ucarı Lake; the Lebanese site authority; "
          "the body managing Gomati on Lemnos",
          "Apply the adapted sheets themselves, without us, in the season "
          "after the project ends", TBD],
         ["Public authorities with territorial competence",
          "North Aegean Periphery; Denizli provincial authorities; the "
          "Lebanese ministry responsible for the coastal site",
          "Adopt or endorse the method as management practice for sites they "
          "are responsible for", TBD],
         ["Tourism operators at the three sites",
          "Beach concessions, boat and tour operators, accommodation and "
          "catering around the three shores",
          "Understand why visitor channelling protects the asset their income "
          "depends on, and cooperate with it", TBD],
         ["Local communities and visitors",
          "Residents of the three territories and seasonal visitors, with the "
          "volunteer network Pi Youth already runs",
          "Recognise the measures on site and not walk round them", TBD],
         ["Other Mediterranean site managers",
          "Coastal and lacustrine managers in the programme area who were "
          "never partners",
          "Find the method, understand whether it fits their shore, and use "
          "it", TBD],
         ["Programme and policy level",
          "Interreg NEXT MED, the MMM database, the AMMIRARE partnership that "
          "produced the original method",
          "See a transfer that worked and a method that has been extended", ""],
         ["Media",
          "Regional press and broadcasters in the three territories; "
          "specialist coastal and tourism press",
          "Carry the story of a shore being managed rather than a project "
          "being funded", TBD]],
        widths=[3.4, 5.0, 5.4, 2.0])
    D.p("Every figure marked [FROM PARTNER] is deliberately blank. Criteria "
        "1.3 and 2.3 both require target groups to be QUANTIFIED, and the "
        "partners hold those numbers: visitor counts at Ucarı Lake, the size "
        "of the volunteer network, the number of operators at each shore. "
        "Inventing them would be worse than leaving them.",
        size=9.5, colour=RED)

    # ------------------------------------------------------- the products
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("3. What we actually produce, and why each one is communication")
    D.p("The strongest communication product of this project is the method "
        "itself, in a form somebody else can pick up. Everything below is "
        "chosen because it carries the method rather than because it carries "
        "the logo.", size=9.5, colour=GREY)

    D.h("3.1 The adapted method pack, in four languages", size=11, space=10)
    D.p("The capitalised sheets, adapted for marine, lacustrine and riparian "
        "sediment shores, as field sheets a site manager can take outdoors: "
        "what to measure, with what, how often, and what the reading means. "
        "English, Greek, Turkish and Arabic. The translation is already a "
        "budgeted line and not an afterthought, because the source material is "
        "in Italian and a method that stays in Italian has not been "
        "transferred.")

    D.h("3.2 The sites themselves", size=11, space=10)
    D.p("The visitor channelling, the boardwalks and the demarcation at the "
        "three demonstration sites are physical communication, seen by every "
        "visitor in the season, and they reach an audience no leaflet reaches. "
        "On site interpretation at each one explains what is being protected "
        "and why the route goes where it goes. Designed for accessibility, "
        "which is also the horizontal principle on disability that criterion "
        "7.1 asks about.")

    D.h("3.3 Shore Custodians", size=11, space=10)
    D.p("MMM_IF01, Beach Custodians, is one of the outputs we capitalise, and "
        "it is simultaneously a capacity activity and a communication channel: "
        "young people doing biodiversity monitoring and dune stewardship on "
        "the sites, who then explain it to visitors. Pi Youth already runs the "
        "volunteer network this builds on.")

    D.h("3.4 Peer to peer transfer, which is the one that produces "
        "replication", size=11, space=10)
    D.p("Site managers listen to other site managers. Once the Turkish and "
        "Lebanese teams have run the method for a season, they present it to "
        "managers of comparable shores in their own countries, in their own "
        "language, as practitioners rather than as a project. That is the "
        "activity most likely to produce use outside the partnership, and it "
        "costs very little.")

    D.h("3.5 Into the MMM database", size=11, space=10)
    D.p("The adapted method, the cross typology comparability protocol and the "
        "lacustrine adaptation are submitted as outputs to the same database "
        "the project drew from, so the next capitalisation call can transfer "
        "them onward. Durable, free, and exactly what the programme built the "
        "database for.")

    D.h("3.6 The conventional channels, kept in proportion", size=11, space=10)
    D.p("The project website provided by the Programme, partner channels, "
        "regional media at each site, and one final event hosted at a "
        "demonstration site rather than in a conference room. These are "
        "necessary and they are not where the points are. They are listed last "
        "on purpose.")

    # ------------------------------------------------------- measurement
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("4. How we will know whether it worked")
    D.p("Communication indicators that measure activity rather than effect are "
        "the reason this criterion is usually weak. Posts published and "
        "leaflets printed measure our effort, not anybody else's behaviour.",
        size=9.5, colour=GREY)
    D.table(["What we measure", "Why it is the right measure", "Target"],
            [["Site managers outside the partnership who request the method "
              "pack", "Demand we did not create by asking", TBD],
             ["Shores outside the three demonstration sites where the method "
              "is applied during the project",
              "Replication, which is criterion 5.3", TBD],
             ["Public authorities that formally adopt or endorse the method",
              "Adoption, which is what makes it outlast the funding", TBD],
             ["Outputs accepted into the MMM database",
              "The method is findable by the next project without us", "3"],
             ["Method pack downloads by language",
              "Shows whether the translation reached the territories it was "
              "for, which a single English total would hide", TBD],
             ["Operators and volunteers trained at each site",
              "Capacity that remains in the territory", TBD]],
            widths=[6.0, 7.0, 2.4])

    # ------------------------------------------------------- greening
    D.h("5. Greening the communication itself")
    D.p("Criterion 5.1 asks about eco-friendly measures in the project's own "
        "daily activities, specifically naming events and communication "
        "material, and the Programme publishes a Guide for greening project "
        "implementation. This is cheap to answer well and most applicants do "
        "not answer it at all.")
    D.bullet("Digital by default. Printed material only where the audience is "
             "physically on a shore and has no other way to receive it.")
    D.bullet("On site interpretation in durable materials appropriate to a "
             "coastal environment, so it is not reprinted every season.")
    D.bullet("Meetings paired with site visits that were happening anyway, "
             "and online where no site is involved, which is also why the "
             "Steering Committee alternates.")
    D.bullet("No single use promotional items.")

    # ------------------------------------------------------- reviewer notes
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("NOTES FOR YOU, DELETE BEFORE FORWARDING", size=14, colour=RED)
    for t, b in [
        ("Why this is worth the effort",
         "Criterion 4.4 counts double. 8 points of 88, the largest single item "
         "on the grid, against a threshold of 75 where only 13 points may be "
         "lost. A generic communication section scores 2 of 4, which doubled "
         "is 4 points gone on its own."),
        ("The blanks are the work",
         "Every [FROM PARTNER] is a number only a partner can give. Acıpayam's "
         "visitor figures have been outstanding since September and they carry "
         "two criteria, not one: the audience table here and the quantified "
         "target groups in 3.1.3. One request to each partner covers both."),
        ("What I would not let a partner add",
         "A social media campaign with follower targets, and a project "
         "newsletter. Both are standard, both measure our effort rather than "
         "anyone else's behaviour, and both make the section read like every "
         "other application. If a partner insists, put them in the activity "
         "list and keep them out of the objectives."),
        ("One thing to check with the Joint Secretariat",
         "The Guidelines say the project website is provided free by the "
         "Programme and is therefore not eligible as a cost, but that managing "
         "it is eligible. Worth confirming what the Programme's site gives us "
         "before any partner budgets for web development, because separate "
         "platforms need the communication manager's approval."),
    ]:
        D.h(t, size=11, colour=RED, space=9)
        D.p(b, size=10)

    out = Path(OUT)
    D.d.save(out)
    n_tbd = sum(1 for t in D.d.tables for r in t.rows for c in r.cells
                if TBD in c.text)
    print(f"wrote {out}")
    print(f"  {len(D.d.paragraphs)} paragraphs, {len(D.d.tables)} tables")
    print(f"  {n_tbd} figures left blank for partners to supply, deliberately")
    # the dash rule must not have eaten anything
    body = "\n".join(p.text for p in D.d.paragraphs)
    assert " , " not in body and not body.strip().startswith(","), \
        "the dash rule has produced a stray comma, check clean()"
    print("  dash rule clean")
    return out


if __name__ == "__main__":
    main()
