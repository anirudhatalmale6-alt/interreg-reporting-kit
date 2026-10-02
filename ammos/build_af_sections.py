#!/usr/bin/env python3
"""Application Form sections you can paste, generated from the budget model.

    python3 build_af_sections.py

YOUR QUESTION: "What do you think we can have automated filled from the
application form? Do you think WP1 and WP2 can be automated?"

THE HONEST SPLIT, SECTION BY SECTION

Three kinds of field in that form, and only two of them should ever be
generated.

  COMPUTABLE. Budget per work package, the Gantt, partner counts, country
  lists, totals. These are arithmetic over data we already hold, they are
  tedious and error-prone by hand, and a script cannot get them wrong twice
  in the same way. Generate all of it.

  FORMULAIC. Work package 1, management and coordination. The form asks for
  the coordinator, the involved partners and the key management tools. Award
  criterion 4.1 is worth 4 points and rewards clarity and named instruments -
  a well-structured conventional answer scores full marks, because there is
  no originality to demonstrate in a steering committee. Generate it, then
  read it.

  JUDGEMENT. Relevance, transnational dimension, target groups, synergies,
  and - this is the one that matters - the COMMUNICATION STRATEGY. Do not
  generate these.

WHY WP2 IS THE EXCEPTION, AND IT IS THE MOST IMPORTANT THING ON THIS PAGE

Criterion 4.4, communication, is scored by the external evaluators at STEP 1
and IT COUNTS DOUBLE: 8 points of the 88, the largest single item on the
entire grid. The threshold is 75 of 88, so we may lose 13 points in total.

Generic communication text scores 2 out of 4. Doubled, that is 4 points gone -
nearly a third of the whole allowance - on the one section where every
applicant writes the same paragraph about social media and a project website.

So WP1 is automated here and WP2 deliberately is not. WP2 is drafted, named
stakeholder by named stakeholder, and it is drafted separately.

WHAT THE FORM'S STRUCTURE REQUIRES THAT OURS DOES NOT YET HAVE

The exported form reveals a three-level hierarchy:

    WORK PACKAGE  ->  OUTPUT  ->  ACTIVITY

Activities are numbered Activity {WP}.{Output}.{Activity}, and each OUTPUT
needs a title, a TARGET VALUE and a SEMESTER OF DELIVERY. Our work package
document numbers activities WP.Activity, two levels, and has no formal output
list with target values at all.

That gap has to be closed by hand before any of this can be pasted, and it is
also award criterion 2.4, indicators, which is 4 more points currently at
risk. Better found now than on 28 October.
"""
import datetime as dt
import pathlib
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt, RGBColor

import pathlib as _pl
sys.path.insert(0, str(_pl.Path(__file__).resolve().parent.parent))
from _shared import here

from build_wp_blended import (DIRECT, P, STAFF, TOTAL_ELIGIBLE, WPS, matrix,
                              SITES)

BASE = str(here(__file__))
OUT = f"{BASE}/AMMOS_AF_sections_to_paste_2026-10-02.docx"

NAVY = RGBColor(0x1F, 0x38, 0x5E)
GREY = RGBColor(0x5A, 0x6B, 0x7D)
RED = RGBColor(0xB4, 0x55, 0x3C)
FLAT = 0.30                       # CC2 + CC3, both 15% of staff
MONTHS = 24


def eur(x):
    """Form convention: 1.300.000,00 €"""
    return f"{x:,.2f}".translate(str.maketrans({",": ".", ".": ","})) + " €"


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Cm(1.8)
        n = self.d.styles["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(10)

    def h(self, t, size=13, colour=NAVY, space=12):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(space)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(t)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), colour

    def p(self, t, size=10, colour=None, italic=False, bold=False):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run(t)
        r.font.size, r.italic, r.bold = Pt(size), italic, bold
        if colour:
            r.font.color.rgb = colour

    def table(self, head, rows, widths=None, right_from=1):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, x in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(x)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), NAVY
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(8.5)
                if str(row[0]).startswith("TOTAL"):
                    r.bold = True
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t


def main():
    elig = {k: DIRECT[k] + STAFF[k] * FLAT for k in DIRECT}
    tot_elig = sum(elig.values())
    tot_direct = sum(DIRECT.values())
    m = matrix()                       # partner -> {WPn: direct cost}
    # WPS entries are TUPLES: (id, title, type, lead, months, desc,
    # activities, roles). Read them by index with names, not by guessing a
    # dict interface that does not exist.
    I_ID, I_TITLE, I_TYPE, I_LEAD, I_MONTHS = 0, 1, 2, 3, 4
    wp_ids = [w[I_ID] for w in WPS]          # "WP1" ... "WP6"

    countries = sorted({P[k][1] for k in P})
    eu = sorted({P[k][1] for k in P if P[k][2] == "EUMC"})
    mpc = sorted({P[k][1] for k in P if P[k][2] == "MPC"})
    n_eu = sum(1 for k in P if P[k][2] == "EUMC")
    n_mpc = sum(1 for k in P if P[k][2] == "MPC")

    D = Doc()
    D.h("AMMOS — Application Form sections ready to paste", size=16)
    D.p(f"Generated {dt.date.today():%d/%m/%Y} from the agreed budget. "
        f"Everything here is computed, not typed. Paste it into the online "
        f"platform and check the platform's own totals agree.",
        size=9.5, colour=GREY, italic=True)

    # ------------------------------------------------- 1. Main information
    D.h("1. Main information — what the platform should show once every "
        "partner is entered")
    D.p("These fields are calculated by the platform from the partners you "
        "add, so you do not type them. But check them against this list after "
        "you have entered everyone: if they disagree, a partner is missing or "
        "has been given the wrong country or role.", size=9.5, colour=GREY)
    D.table(["Field", "Should read", "Currently reads in your export"],
            [["Title / Acronym", "AMMOS", "AMMOS"],
             ["Name of the Applicant", "Hellenic Centre for Marine Research",
              "Hellenic Centre for Marine Research"],
             ["Duration (months)", str(MONTHS), "24"],
             ["Type of Project", "2, Green Tourism projects",
              "2, Green Tourism projects"],
             ["Programme Specific Objective",
              "2.2 (RSO2.4) Promoting climate change adaptation",
              "2.2 (RSO2.4) ..."],
             ["EU Partners / Countries",
              f"{n_eu} partners / {len(eu)} countries ({', '.join(eu)})",
              "1 / — NOT YET ENTERED"],
             ["MPC Partners / Countries",
              f"{n_mpc} partners / {len(mpc)} countries ({', '.join(mpc)})",
              "0 — NOT YET ENTERED"],
             ["Associated partners",
              "1 once North Aegean signs, 2 if Sant'Anna also accepts", "0"],
             ["International Organizations", "0", "0"],
             ["Geographic coverage", ", ".join(countries),
              "Greece — INCOMPLETE"],
             ["Budget", eur(tot_elig), "blank"]],
            widths=[5.0, 7.5, 5.0])

    # ------------------------------------- 4.2 budget per work package
    D.h("4.2 Budget per Work Package — fully computed, paste as is")
    D.p("NOTE THE DENOMINATOR. This table is in DIRECT costs, not total "
        "eligible costs: the form's own heading is \"Total budgeted direct "
        "cost per PP\", and the two rows beneath it split those direct costs "
        "between EU Mediterranean and Mediterranean Partner Countries. So the "
        "50% figure the platform will display is the DIRECT one. Both are "
        "above the floor, but they are different numbers and you should not "
        "be surprised when the platform shows the second.",
        size=9.5, colour=RED)
    head = ["Partner"] + wp_ids + ["TOTAL direct"]
    rows = []
    for k in sorted(P, key=lambda k: -DIRECT[k]):
        rows.append([P[k][0]] + [eur(m[k].get(w, 0)) for w in wp_ids]
                    + [eur(DIRECT[k])])
    rows.append(["TOTAL per WP"]
                + [eur(sum(m[k].get(w, 0) for k in P)) for w in wp_ids]
                + [eur(tot_direct)])
    D.table(head, rows, widths=[4.2] + [1.9] * len(wp_ids) + [2.3])

    mpc_direct = sum(DIRECT[k] for k in P if P[k][2] == "MPC")
    eu_direct = tot_direct - mpc_direct
    D.table(["", "Direct cost", "% of total direct"],
            [["Total direct costs EU Mediterranean Countries", eur(eu_direct),
              f"{eu_direct/tot_direct:.2%}".replace(".", ",")],
             ["Total direct costs Mediterranean Partner Countries",
              eur(mpc_direct), f"{mpc_direct/tot_direct:.2%}".replace(".", ",")],
             ["TOTAL", eur(tot_direct), "100,00%"]],
            widths=[9.0, 4.0, 4.0])
    D.p(f"For your own check, not for the form: on TOTAL ELIGIBLE costs — the "
        f"denominator Guidelines 4.4.3 actually names — the MPC share is "
        f"{sum(elig[k] for k in P if P[k][2]=='MPC')/tot_elig:.2%}"
        .replace(".", ",") +
        f". Both are above 50%.", size=9.5, colour=GREY)

    # ------------------------------------------------ 4.4 the Gantt
    D.h("4.4 Work plan timeline — one row per work package, M1 to M24")
    D.p("Shaded where the work package runs. Transfer the pattern into the "
        "platform's grid.", size=9.5, colour=GREY)
    gh = ["WP"] + [f"M{i}" for i in range(1, MONTHS + 1)]
    grows = []
    import re as _re
    for w in WPS:
        # "M1-M24" or "M1–M24" - the dash may be an en dash, which is why this
        # pulls the digits out rather than splitting on a character.
        nums = [int(x) for x in _re.findall(r"M(\d+)", w[I_MONTHS])]
        a, b = (nums + [1, MONTHS])[:2] if len(nums) >= 2 else (1, MONTHS)
        assert 1 <= a <= b <= MONTHS, (
            f"{w[I_ID]} runs {w[I_MONTHS]}, which is outside M1 to M{MONTHS}")
        grows.append([w[I_ID]]
                     + ["X" if a <= i <= b else "" for i in range(1, MONTHS + 1)])
    D.table(gh, grows, widths=[1.3] + [0.62] * MONTHS)

    # ------------------------------------------------------- WP1, generated
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("WP1 Management and coordination — generated, read it before pasting")
    D.p("WP coordinator: Applicant (HCMR). The form fixes both mandatory work "
        "packages to the Applicant and FAQ 1.11 says the same, so this is not "
        "a choice.", size=9.5, colour=GREY)
    D.p("Involved partners: " + ", ".join(P[k][0] for k in P), size=9.5,
        colour=GREY)
    D.h("Description of key management tools", size=11, space=8)
    for t in [
        "Steering Committee. One voting representative per partner, chaired "
        "by the Lead Partner, meeting at the kick-off, then every six months, "
        "four times in total, alternating between online and the demonstration "
        "territories so that each meeting is paired with a site visit. It "
        "approves the work plan, the reallocation of tasks and any "
        "modification request before it is submitted.",
        "Project Management Team. The Lead Partner's project manager and "
        "financial officer, meeting monthly online with the work package "
        "leaders. It is the body that notices slippage, which is why it meets "
        "between Steering Committees rather than alongside them.",
        "Shared procurement and deliverables tracker. One file per programme "
        "on a shared workspace, one row per contract position, carrying the "
        "planned and actual publication, offer deadline and signature dates, "
        "the status and the days late. It is updated by the partner that owns "
        "the contract, not by the Lead Partner, so the record is maintained by "
        "whoever holds the facts.",
        "Risk register. Reviewed at every Steering Committee, with a named "
        "owner and a mitigation for each entry. It opens with the risks "
        "already identified in this proposal rather than being created empty: "
        "the security situation affecting travel to one demonstration site, "
        "the dependency of three national baselines on a common parameter set, "
        "and the seasonality of visitor-pressure data collection.",
        "Quality assurance of deliverables. Every deliverable is reviewed by a "
        "partner that did not produce it before it is accepted, and the "
        "territorial baselines are additionally reviewed against the "
        "capitalised method by the partner that holds that expertise.",
        "Reporting cycle. Partner financial and activity reports to the Lead "
        "Partner, consolidated into the Programme's reporting periods, with an "
        "internal deadline set ahead of the Programme's own so that a late "
        "partner does not make the consortium late.",
    ]:
        D.p("• " + t, size=10)

    D.h("Activities under WP1", size=11, space=8)
    D.table(["Activity", "Title", "Months"],
            [["1.1.1", "Project start-up: partnership agreement, internal "
                       "procedures, kick-off Steering Committee", "M1-M3"],
             ["1.1.2", "Ongoing coordination, monitoring and risk management",
              "M1-M24"],
             ["1.1.3", "Financial management, reporting and verification of "
                       "expenditure", "M1-M24"]],
            widths=[2.0, 10.5, 2.5])

    # ------------------------------------------------------- WP2, NOT generated
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("WP2 Communication and dissemination — DELIBERATELY NOT GENERATED",
        colour=RED)
    D.p("This is the one section on the form that must not be filled with "
        "competent generic text, and the arithmetic says so.", bold=True)
    D.p("Criterion 4.4 is scored by the external evaluators at Step 1 and it "
        "COUNTS DOUBLE: 8 points of 88, the largest single item on the grid. "
        "The threshold is 75 of 88, so the whole proposal may lose 13 points. "
        "A generic communication section scores 2 out of 4, which doubled is 4 "
        "points — nearly a third of the entire allowance, lost on the section "
        "where every applicant writes the same paragraph about a website and "
        "social media.", size=10)
    D.p("It is drafted separately, by stakeholder, with named channels and "
        "quantified reach. What the form asks for here is the objectives of "
        "the communication strategy and the approach to communication and "
        "visibility, and both answers have to be specific to sustainable "
        "tourism and to these territories.", size=10)

    out = pathlib.Path(OUT)
    D.d.save(out)
    print(f"wrote {out}")
    print(f"  budget per WP: {len(P)} partners x {len(wp_ids)} work packages, "
          f"direct total {tot_direct:,.2f}")
    print(f"  EU direct {eu_direct:,.2f} ({eu_direct/tot_direct:.2%}) | "
          f"MPC direct {mpc_direct:,.2f} ({mpc_direct/tot_direct:.2%})")
    # the per-WP matrix must reconcile to each partner's direct cost
    for k in P:
        s = sum(m[k].values())
        assert abs(s - DIRECT[k]) < 0.01, (
            f"{k}: work package shares sum to {s:,.2f} but its direct cost is "
            f"{DIRECT[k]:,.2f}. The form would not reconcile.")
    assert abs(sum(sum(m[k].values()) for k in P) - tot_direct) < 0.01
    print("  reconciled: every partner's WP shares sum to its direct cost")
    return out


if __name__ == "__main__":
    main()
