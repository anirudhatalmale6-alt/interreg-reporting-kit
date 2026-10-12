#!/usr/bin/env python3
"""Partnership texts: 3.3.7 roles, HCMR's 3.3.3, six Key expertise lines, and
the AMMOS half of three partner capacity boxes.

    python3 build_partnership_texts.py

He sent the counters: "Key expertise provided to the project" 100,
"3.3.3 / 3.3.4 Experience and capacity" 2000. Same pair for the Applicant and
for every partner.

AND A GAP I HAD MISSED. Box 3.3.7, "Roles and tasks", is empty in every export
he has sent - the instruction text runs straight into 3.4.1. Nobody wrote it,
including me. This morning I told him the capacity boxes would be written
"consistent with the roles already described in 3.3.7". They are not described
there; they are described inside 3.4.1. That sentence was false and he is
told so.

WHAT EVERY CLAIM BELOW RESTS ON

  HCMR facts       Wikipedia's HCMR article, read 12 Oct (hcmr.gr itself
                   returned HTTP 500): formed 3 June 2003 from the National
                   Centre for Marine Research and the Institute of Marine
                   Biology of Crete; under the General Secretariat for Research
                   and Innovation; HQ Anavyssos; three institutes, one of them
                   Marine Biological Resources and INLAND WATERS; Hydrobiological
                   Station of Rhodes; research vessels; publishes Mediterranean
                   Marine Science.
  roles and leads  the application itself - 3.4.1 "expertise is matched to
                   task", 3.4.2 WP leads (WP3 HCMR, WP4 Aegean, WP5 PIYA, WP6
                   LCEC), and the CNR authorship statement in 3.2.1.
  associate        statutory remit only, as in yesterday's draft.

NOT CLAIMED: any count of past EU projects, staff numbers, or named projects.
Those are the partner's to supply. The texts are complete without them - there
is no slot, no bracket and no XXXX to paste into the platform by accident,
because the last three exports each carried a placeholder that went live.
"""
import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from _shared import here

BASE = str(here(__file__))
OUT = f"{BASE}/AMMOS_Partnership-texts_3.3.3-3.3.4-3.3.7_2026-10-12.docx"

NAVY = RGBColor(0x1F, 0x38, 0x5E)
GREY = RGBColor(0x5A, 0x6B, 0x7D)
RED = RGBColor(0xB4, 0x55, 0x3C)

LIM_KEY, LIM_CAP, LIM_ROLES = 100, 2000, 2000
FORBIDDEN = ("elke", "finance office", "non-compliant", "option b", "ras ",
             "wrong", "4.4.3", "remaco", "887", "tbd", "tbc", "xxx", "[", "]",
             "???")

# Each WP lead as the application states it - the texts are checked against it
LEADS = {"HCMR": ["WP1", "WP2", "WP3"], "Aegean": ["WP4"], "PIYA": ["WP5"],
         "LCEC": ["WP6"], "CNR": [], "Acipayam": []}

KEY = {
    "HCMR (Applicant)":
        "Scientific coordination; marine and inland-water ecology; biological "
        "and ecological method sheets",
    "University of the Aegean":
        "Demonstration design on Aegean island shores; tourism baseline, "
        "visitor survey and valuation",
    "Pi Youth Association":
        "Youth engagement, volunteering and school custodian groups; youth "
        "channels and social media",
    "Acipayam Municipality":
        "Authority for the Ucari pond site; municipal site management and "
        "site-manager training",
    "ISAC-CNR":
        "Co-author of the source method; topographic survey and meteo-marine "
        "forcings sheets",
    "LCEC":
        "National public agency: institutional uptake, policy embedding and "
        "management plan adoption",
}

HCMR_333 = [
    "HCMR is Greece's national public centre for marine research, formed in "
    "2003 from the National Centre for Marine Research and the Institute of "
    "Marine Biology of Crete, under the General Secretariat for Research and "
    "Innovation. Its three institutes - Oceanography; Marine Biological "
    "Resources and Inland Waters; and Marine Biology, Biotechnology and "
    "Aquaculture - cover the full range this project needs.",

    "WHY HCMR. AMMOS moves one method across unlike shores: a marine dune "
    "system, a marine coastal site and an inland freshwater pond. Few "
    "institutions work on both sides of that line. HCMR's Institute of Marine "
    "Biological Resources and Inland Waters does, which is why HCMR carries "
    "the biological and ecological method sheets and can judge, sheet by "
    "sheet, where a procedure written for a western Mediterranean beach still "
    "holds on a lake shore.",

    "CAPACITY TO COORDINATE. As Applicant HCMR runs WP1 and WP2, which cannot "
    "be delegated, and leads WP3, the work package that gates everything "
    "after it. It provides the scientific coordinator, who signs cross-site "
    "consistency on every deliverable, and the project manager. As a public "
    "research body it operates the financial, legal and procurement services "
    "that EU-funded projects require.",

    "INFRASTRUCTURE. Field and laboratory capacity sits at its Anavyssos "
    "headquarters in Attica, its facilities in Crete and the Hydrobiological "
    "Station of Rhodes, with research vessels for nearshore work. HCMR also "
    "publishes the journal Mediterranean Marine Science, a direct route for "
    "the adapted method into the regional scientific record.",
]

ROLES_337 = [
    "Each partner leads what it would defend in its own name, and every lead "
    "follows an institutional mandate rather than a budget share.",

    "HCMR, APPLICANT. Coordinates WP1 and WP2 and leads WP3: the transfer "
    "workshop, the 21 transferability decisions and the three baselines. "
    "Provides the scientific coordinator and carries the biological and "
    "ecological method sheets.",

    "ISAC-CNR. The institute behind the source method; its researcher is "
    "MMM_IF02's first named author. Leads adaptation of the topographic and "
    "forcings sheets in WP3, so the method is adapted by people who know where "
    "it was designed to stop.",

    "UNIVERSITY OF THE AEGEAN. Leads WP4: site designs and consents, "
    "installation and one monitored visitor season, and the tourism and "
    "valuation baseline at Gomati on Lemnos, in its own region.",

    "PI YOUTH ASSOCIATION. Leads WP5: custodian material, nine school groups, "
    "youth channels and social media - its own field, not a side activity.",

    "ACIPAYAM MUNICIPALITY. The authority for the Ucari pond site, which it "
    "will run after the project. Leads site-manager training in WP5 and takes "
    "the adoption decision for its own plan in WP6.",

    "LCEC. Leads WP6: costed management plans, adoption decisions, "
    "stewardship committees, the transferability guide and policy "
    "recommendations - the institutional work a national public agency does "
    "well.",

    "NORTH AEGEAN WATER DIRECTORATE, ASSOCIATED. Holds the statutory water and "
    "flood-risk planning remit for the North Aegean, giving the Gomati results "
    "a route into the plans that govern the island's coast after closure.",

    "WHY THIS DISTRIBUTION. Each receiving territory has a partner with "
    "authority over its site or its planning, the method's authors sit inside "
    "the partnership, and the four technical work packages are led by four "
    "different organisations.",
]

# The AMMOS half of a 3.3.4 box. The partner writes its own record ABOVE this.
ROLE_334 = {
    "University of the Aegean": (
        "IN AMMOS the University of the Aegean leads WP4, demonstration. Based "
        "in the North Aegean, the region of the Gomati site on Lemnos, it leads "
        "site design and consents, installation and one full monitored visitor "
        "season, and the tourism and valuation baseline: visitor numbers and "
        "seasonality, willingness to visit and to pay, and the economic value "
        "of the shore to the local economy. Its Department of Marine Sciences "
        "works on the coasts of the same island region, which is what a "
        "demonstration on Lemnos needs."),
    "ISAC-CNR": (
        "IN AMMOS ISAC-CNR brings the source method itself. CNR was a partner "
        "in AMMIRARE, and its researcher is the first named author of MMM_IF02, "
        "the guidelines AMMOS transfers. It leads adaptation of the "
        "topographic and meteo-marine forcings sheets in WP3, where knowing why "
        "each procedure was designed as it was decides whether it survives the "
        "move to a dune system, a Lebanese coast and an inland pond."),
    "Acipayam Municipality": (
        "IN AMMOS Acipayam Municipality is the authority for the Ucari "
        "recreation pond, the project's inland freshwater site, and the body "
        "that will manage it after the project ends. It leads site-manager "
        "training in WP5, so the staff who will run the site use the method "
        "before they are asked to adopt it, and it takes the adoption decision "
        "for the Ucari management plan in WP6."),
}


def clean(t):
    return re.sub(r"\s*[–—]\s*", ", ", str(t))


def joined(paras):
    return "\n\n".join(clean(p) for p in paras)


def check():
    fails = []
    for k, v in KEY.items():
        if len(v) > LIM_KEY:
            fails.append(f"Key expertise {k}: {len(v)} > {LIM_KEY} - cut "
                         f"{len(v) - LIM_KEY}")
    h = joined(HCMR_333)
    if len(h) > LIM_CAP:
        fails.append(f"3.3.3: {len(h)} > {LIM_CAP} - cut {len(h) - LIM_CAP}")
    r = joined(ROLES_337)
    if len(r) > LIM_ROLES:
        fails.append(f"3.3.7: {len(r)} > {LIM_ROLES} - cut "
                     f"{len(r) - LIM_ROLES}")
    for who in ("HCMR", "ISAC-CNR", "AEGEAN", "PI YOUTH", "ACIPAYAM", "LCEC",
                "NORTH AEGEAN"):
        if who not in r:
            fails.append(f"3.3.7 never mentions {who}")
    for org, wps in LEADS.items():
        token = {"PIYA": "PI YOUTH", "Aegean": "UNIVERSITY OF THE AEGEAN",
                 "CNR": "ISAC-CNR"}.get(org, org).upper()
        para = next((p for p in ROLES_337 if p.upper().startswith(token)
                     or token in p.upper()[:40]), "")
        if not para:
            fails.append(f"3.3.7: no paragraph found for {org}")
        for wp in wps:
            if f"Leads {wp}" not in para and f"leads {wp}" not in para \
                    and wp not in para:
                fails.append(f"3.3.7: {org} leads {wp} but its paragraph "
                             f"does not say so")
    if "four technical work packages are led by four different" not in r:
        fails.append("the distribution claim was changed - re-verify it "
                     "against the WP leads before shipping")
    for k, v in ROLE_334.items():
        if len(v) > 900:
            fails.append(f"{k} role paragraph {len(v)} chars - leaves under "
                         f"1100 for the partner's own record")
    every = (list(KEY.values()) + HCMR_333 + ROLES_337
             + list(ROLE_334.values()))
    for t in every:
        low = t.lower()
        for bad in FORBIDDEN:
            if bad in low:
                fails.append(f"forbidden or placeholder {bad!r} in: "
                             f"{t[:50]}")
    return fails


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Cm(1.5)
        n = self.d.styles["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(10)

    def h(self, t, size=13, colour=NAVY, space=12):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(space)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(clean(t))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), colour

    def p(self, t, size=10, colour=None, italic=False, bold=False):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run(clean(t))
        r.font.size, r.italic, r.bold = Pt(size), italic, bold
        if colour:
            r.font.color.rgb = colour

    def table(self, head, rows, widths=None, small=8.5):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, x in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(clean(x))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(small), NAVY
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run(clean(v)).font.size = Pt(small)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)


def main():
    fails = check()
    assert not fails, "NOT SHIPPING:\n  - " + "\n  - ".join(fails)
    h, r = joined(HCMR_333), joined(ROLES_337)

    D = Doc()
    D.h("AMMOS - Partnership texts: 3.3.7, 3.3.3, Key expertise and the "
        "partner capacity boxes", size=15)
    D.p(f"Generated {dt.date.today():%d/%m/%Y}. Limits from your screenshots: "
        f"Key expertise 100, capacity boxes 2000. Nothing below contains a "
        f"bracket, a blank or an XXXX - every text is complete as it stands.",
        colour=GREY, italic=True)

    D.h("First, a correction: 3.3.7 Roles and tasks is EMPTY", size=12,
        colour=RED)
    D.p("In every export you have sent, the 3.3.7 question runs straight into "
        "3.4.1 with no answer. Nobody has written it - including me. This "
        "morning I told you the capacity boxes would match \"the roles "
        "already described in 3.3.7\". They are not there; they are inside "
        "3.4.1. That sentence was wrong. 3.3.7 is the box an assessor reads "
        "to judge whether the distribution of roles makes sense, so it is "
        "the first text below.")

    D.h(f"1. 3.3.7 Roles and tasks - {len(r)} characters", size=12)
    for p in ROLES_337:
        D.p(p)
    D.p("I have not seen the counter on 3.3.7, so it is written to 2000, the "
        "size of the boxes around it. If it turns out larger, send the "
        "number and I will extend it rather than leave it half-filled.",
        size=8.5, colour=GREY, italic=True)

    D.h(f"2. HCMR - 3.3.3 Experience and capacity of the Applicant - "
        f"{len(h)} of 2000", size=12)
    for p in HCMR_333:
        D.p(p)
    D.p(f"About {2000 - len(h)} characters are left on purpose. The one thing "
        f"this box should still say, and only HCMR can supply, is two or "
        f"three EU-funded projects HCMR has coordinated or taken part in, "
        f"ideally coastal. One sentence at the end. I have not invented any.",
        size=8.5, colour=GREY, italic=True)
    D.p("Every institutional fact here comes from HCMR's public record, read "
        "today: the 2003 merger, the three institutes, Anavyssos, Rhodes, the "
        "journal. Please still read it once with HCMR's eyes - it is your "
        "own institution.", size=8.5, colour=GREY, italic=True)

    D.h("3. Key expertise provided to the project - 100 characters, all six",
        size=12)
    D.table(["Organisation", "Text", "Chars"],
            [[k, v, str(len(v))] for k, v in KEY.items()],
            widths=[4.2, 12, 1.3])

    D.h("4. The AMMOS half of the partner capacity boxes (3.3.4)", size=12)
    D.p("Your email asks each partner for its background and past projects. "
        "That is the right half for them to write. The other half - what "
        "they do in AMMOS - has to match 3.3.7 and the work packages word for "
        "word, and partners writing it independently is how boxes end up "
        "contradicting each other. So each partner gets this paragraph to put "
        "at the END of its box, with its own record above it. Total must stay "
        "under 2000.")
    for k, v in ROLE_334.items():
        D.h(f"{k} - {len(v)} characters, leaves {2000 - len(v) - 2} for "
            f"their own record", size=10.5, space=8)
        D.p(v)
    D.p("Pi Youth Association and LCEC already have their boxes filled. PIYA "
        "only needs the duplicated paragraphs removed.", colour=GREY)

    D.d.save(OUT)
    print(f"wrote {OUT}")
    print(f"  3.3.7 {len(r)}  3.3.3 {len(h)}  key {[len(v) for v in KEY.values()]}"
          f"  3.3.4 {[len(v) for v in ROLE_334.values()]}")


if __name__ == "__main__":
    main()
