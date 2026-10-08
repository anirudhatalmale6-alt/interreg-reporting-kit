#!/usr/bin/env python3
"""Checklist B Sections 2 and 3 - thirty-one boxes.

    python3 build_checklist_b_2_3.py

Section 2 is a LOCATION grid: eleven categories, each YES / NO / N/A with a
500-character Specify box. Section 3 is the significance test: nineteen
questions plus Other comments, all 500.

THE POSTURE, AND IT IS THE SAME ONE AS CHECKLIST A

Answer YES wherever a site genuinely touches a category, then bound it
precisely in the Specify box. Two reasons.

First, our own application already says it. Section 3.1.1, which is in the
form, states that the Gomati site "sits outside Natura 2000 site GR4110006,
which covers almost the whole coastal and marine perimeter of the island".
Ticking NO against "buffer zone of natural protected area" would contradict
text an assessor can read two tabs away. Internal inconsistency is worse than
any individual YES.

Second, the whole argument of this project is that a habitat under pressure
which falls between two designations is exactly the case a transferable
management method is for. Denying the proximity throws away our own premise.

So: coastal YES, marine YES, buffer zone YES, vulnerable landscape YES, and
each one explained and bounded. Densely populated NO, estuaries NO, cultural
heritage NO - and those are NO because they are false, not because NO is
comfortable.

SECTION 3 IS A SIGNIFICANCE TEST AND OUR ANSWERS ARE CONSISTENT

Nineteen questions that all ask, in different words, how big and how permanent
the effect is. The honest answers run one way throughout: small, local,
temporary, intermittent, reversible, avoidable. Question 18 - would the impact
be irreversible - is the one the whole design was built to answer, and the
answer is no, with the reason already budgeted and already written into three
other sections.
"""
import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from _shared import here

BASE = str(here(__file__))
OUT = f"{BASE}/AMMOS_Checklist-B_Sections-2-and-3_2026-10-10.docx"

NAVY = RGBColor(0x1F, 0x38, 0x5E)
GREY = RGBColor(0x5A, 0x6B, 0x7D)
RED = RGBColor(0xB4, 0x55, 0x3C)
GREEN = RGBColor(0x2E, 0x6B, 0x45)
AMBER = RGBColor(0x9A, 0x6B, 0x1E)

LIM = 500
RANGE = re.compile(r"(?<=\d)\s*[–—]\s*(?=[A-Za-z]?\d)")


def clean(t):
    t = str(t)
    t = RANGE.sub("\x00", t)
    t = re.sub(r"(?<!-)\s*--\s*(?!-)", ", ", t)
    t = re.sub(r"\s*[–—]\s*", ", ", t)
    t = t.replace("\x00", "-")
    return re.sub(r",\s*,", ",", t)


# ------------------------------------------------- Section 2: the location grid
S2 = [
    ("Densely populated area", "NO",
     "None of the three sites is in a densely populated area. Gomati is a "
     "rural shore on Lemnos; Ucari is a recreation pond outside Acipayam in "
     "Denizli province; the Lebanese site is a coastal location outside an "
     "urban centre. Visitor pressure at all three is seasonal and recreational "
     "rather than residential, which is precisely the pressure the project "
     "addresses."),

    ("Cultural heritage site", "NO",
     "No site is a designated cultural heritage site and no intervention "
     "touches one. The Faraklou Geological Park lies adjacent to the Gomati "
     "area on Lemnos, but it is a geological designation rather than a "
     "cultural one and no project works fall inside it. Should any "
     "archaeological feature be identified during site preparation, works "
     "stop and the national procedure applies - all three sites pass their "
     "own national consent regime before installation."),

    ("Natural protected area", "NO",
     "No intervention takes place inside a designated protected area, and "
     "this is central to the project rather than incidental. The Gomati shore "
     "is an EU Habitats Directive Annex I habitat type that lies OUTSIDE "
     "Natura 2000 site GR4110006 and outside the adjacent Faraklou "
     "Geological Park - a habitat under recorded visitor pressure that "
     "belongs to neither designation. That gap is the reason a transferable "
     "management method, rather than a new designation, is the realistic "
     "instrument here."),

    ("Wetlands", "NO",
     "No site is a designated wetland and none is listed under the Ramsar "
     "Convention. Ucari pond in Denizli province is an artificial, "
     "groundwater-fed recreation water body with no outlet to the sea, "
     "constructed for amenity use rather than a natural wetland system. Its "
     "littoral vegetation is nonetheless characterised and monitored, and the "
     "substituted freshwater littoral protocol exists precisely because the "
     "source method's dune-vegetation sheet does not apply there."),

    ("Coastal areas", "YES",
     "Two of the three sites are coastal: the Gomati dune and beach system on "
     "Lemnos, Greece, and the Lebanese coastal site. Both are Mediterranean "
     "sediment shores under seasonal visitor pressure, which is the condition "
     "the transferred method was written for. All interventions at both sites "
     "are demountable and removable without lasting trace, set back from the "
     "active shoreline, and designed from each site's own topographic and "
     "sediment baseline."),

    ("Marine areas", "YES",
     "The two coastal sites include a marine component. The method sheets "
     "AMMOS transfers cover submerged topography and Posidonia meadow "
     "condition, so characterisation and monitoring extend into the nearshore "
     "zone at Gomati and at the Lebanese site. That work is survey and "
     "observation only: no dredging, no bottom or underwater cleaning, no "
     "sediment extraction or importation, and no installation below the "
     "waterline at any site."),

    ("Estuaries", "NO",
     "No site is an estuary or within an estuarine system. Neither coastal "
     "site is at a river mouth, and Ucari is a closed artificial pond with no "
     "outlet. No estuarine hydrology or habitat is affected."),

    ("Mountain areas", "N/A",
     "No site is in a mountain area as the category is normally applied. "
     "Ucari pond sits on the inland plateau of Denizli province near the "
     "Acipayam canyon, which gives it an upland rather than montane setting, "
     "and the project documents that canyon baseline as a worked example of a "
     "site type the source method never addressed. No works take place on "
     "slopes or at altitude."),

    ("Buffer zone of natural protected area", "YES",
     "The Gomati site adjoins Natura 2000 site GR4110006, which covers almost "
     "the whole coastal and marine perimeter of Lemnos, and lies close to the "
     "Faraklou Geological Park. It is outside both, but functionally adjacent, "
     "so it is treated as a buffer situation throughout: every component "
     "demountable and removable without lasting trace, nothing extractive, "
     "design following site characterisation rather than a template, and "
     "national consent and screening before works."),

    ("Special area for biodiversity protection", "NO",
     "No intervention is inside a special area for biodiversity protection. "
     "The Gomati shore is an Annex I habitat type outside the Natura 2000 "
     "perimeter, so it carries habitat value without the designation - which "
     "is why the project produces the condition data any future designation or "
     "management plan would need. Biodiversity state and an ecological status "
     "indicator are characterised at all three sites before any works."),

    ("Vulnerable landscape area", "YES",
     "All three sites are vulnerable landscapes in the sense this project "
     "exists to address: unconsolidated sediment surfaces degrading under "
     "concentrated seasonal foot traffic. The vulnerability is the reason for "
     "the intervention, not a consequence of it. The response is visitor "
     "channelling rather than exclusion - boardwalks and marked routes that "
     "move feet off the vulnerable surface while the site stays publicly "
     "accessible."),
]

# ------------------------------------------- Section 3: the significance test
S3 = [
    (1, "May the intervention cause a change on the environmental "
        "conditions?",
     "Yes, and the intended change is an improvement: less trampling and "
     "erosion across the sediment surface, achieved by concentrating foot "
     "traffic onto a defined route. The unintended changes are localised and "
     "short-lived - disturbance during installation and post footings in the "
     "sediment. Both are measured, because each site is characterised before "
     "works and monitored through one full visitor season after."),

    (2, "Would there be any out-of-scale feature affecting the environmental "
        "context?",
     "No. The built elements are footpath-width boardwalk runs, boundary "
     "marker poles and interpretation signage - small, low, and modular. "
     "Nothing is engineered coastal defence, nothing rises above eye level, "
     "and every component is demountable. Scale was constrained deliberately: "
     "the budget carries no Infrastructures and works category at all, and "
     "fabrication is procured as a service."),

    (3, "Would the area be visually affected?",
     "Modestly and locally. A boardwalk, marker poles and interpretation "
     "panels are visible within the site. Materials and alignment are chosen "
     "at design stage from each site's own baseline, and the design is signed "
     "off by the authority that will own it. Interpretation panels are a "
     "deliberate visual addition: a boundary nobody can see or understand is "
     "a boundary people walk around."),

    (4, "Would such visual effect extend over a large area?",
     "No. The visual effect is confined to the immediate shore surface at each "
     "of the three sites and to the approach path. There is no change to "
     "skyline, ridgeline or seaward view, no lighting, and no structure of a "
     "height visible from outside the site. The affected footprint at each "
     "site is a corridor, not an area."),

    (5, "May the project produce a potential transnational impact on the "
        "environment? To which extent and scale?",
     "Not as a physical effect - the three sites are independent, far apart, "
     "and each intervention is local. The transnational effect is "
     "methodological and positive: a monitoring protocol and a transferability "
     "assessment usable on other Mediterranean sediment shores, returned to "
     "the MMM common database under open licence. The intended scale of that "
     "effect is basin-wide; the scale of any physical effect is three "
     "footpath corridors."),

    (6, "How many people would be affected by the project intervention? "
        "(estimation)",
     "Positively and directly: around 60 site managers and local authority "
     "staff trained, nine school custodian groups across three countries, and "
     "the visitors to the three sites, whose numbers are being measured rather "
     "than estimated - the WP3 visitor survey establishes the real seasonal "
     "figure at each site in the first season. Nobody is displaced, no access "
     "is closed and no livelihood is restricted."),

    (7, "How many other types of ecosystemic subjects (fauna and flora, "
        "businesses, infrastructures) would be affected?",
     "Flora: the dune and littoral vegetation communities at the three sites, "
     "affected positively through reduced trampling. Fauna: the associated "
     "invertebrate and bird communities, affected only through reduced "
     "disturbance. Businesses: local tourism operators, consulted on the "
     "design and affected positively, since the shore stays open. "
     "Infrastructures: none - no utilities, roads or services are altered at "
     "any site."),

    (8, "Would valuable environmental features or resources be affected?",
     "Yes, and the project exists to protect them: Annex I habitat type dune "
     "and beach vegetation at Gomati, nearshore Posidonia at the two marine "
     "sites, and the littoral vegetation of Ucari pond. Each is characterised "
     "before works and monitored after. No feature is removed, cleared or "
     "extracted, and the installations are set to carry traffic over the "
     "vulnerable surface rather than through it."),

    (9, "Is there a risk that any environmental standards may be breached?",
     "No standard is approached, let alone breached. There are no emissions "
     "to air or water, no discharge, no waste stream beyond ordinary "
     "construction packaging, no noise beyond short installation periods and "
     "no use of hazardous materials. Each site passes its own national consent "
     "and environmental screening before works, which is the control that "
     "would catch any site-specific standard we have not anticipated."),

    (10, "Is there a risk that protected sites, areas, features may be "
         "affected?",
     "Gomati adjoins Natura 2000 site GR4110006 without lying inside it, so "
     "the risk is proximity rather than direct effect. It is managed by the "
     "same four constraints that govern everything: reversibility, "
     "characterisation before installation, nothing extractive, and national "
     "consent first. The Annex I habitat type at the site is itself the "
     "feature the project protects, and its condition is measured before and "
     "after."),

    (11, "What is the likelihood that the foreseen effects will actually take "
         "place?",
     "The positive effect - reduced trampling on the vulnerable surface - is "
     "likely but not assumed, which is why it is measured rather than claimed. "
     "The negative effects are certain but trivial: installation disturbance "
     "and post footings will definitely occur, at a known and small scale. "
     "Where the monitoring shows the intended effect is not materialising, the "
     "installation is altered, because it was built to be."),

    (12, "Would the effects continue for a long time? How long? (estimation)",
     "The positive effect is intended to continue indefinitely, sustained by "
     "the adopted management plan and the stewardship committee at each site. "
     "The negative effects do not: installation disturbance lasts days per "
     "site, and the physical presence of the structures lasts only as long as "
     "they are wanted, since they are demountable. No effect is locked in "
     "beyond the life of a removable structure."),

    (13, "Would the effect be cumulative?",
     "No cumulative negative effect is expected. The three sites are "
     "independent and far apart, so effects do not aggregate geographically, "
     "and no other development is planned at any of them during the project. "
     "The positive effect is intended to be cumulative in the useful sense: "
     "monitoring continues after closure through trained staff and the "
     "custodian groups, so the condition record builds rather than stopping."),

    (14, "Would the effect be direct rather than indirect?",
     "Both, and the direction differs. The direct effects are physical and "
     "small: the installation footprint and its disturbance. The main "
     "indirect effect is the larger one and it is positive - redistributing "
     "foot traffic reduces pressure across the wider sediment surface, which "
     "is the mechanism the whole method relies on. A further indirect effect "
     "is institutional: three authorities adopt a costed management plan."),

    (15, "Would the effect be permanent rather than temporary?",
     "Temporary by design. Every component is demountable and removable "
     "without lasting trace - there is no permanent foundation, no hard "
     "surfacing and no excavation beyond post footings. Reversibility is "
     "stated in the work package, in section 3.5.1 on environmental "
     "sustainability and in the do-no-significant-harm answer, and it is why "
     "the budget carries nothing in the Infrastructures and works category."),

    (16, "Would the impact be continuous rather than intermittent?",
     "Intermittent, and seasonal. Installation work is confined to short "
     "defined periods at each site. Visitor use, and therefore the pressure "
     "the project manages, concentrates in the summer season at all three "
     "sites, which is also why monitoring is organised around one full visitor "
     "season rather than a continuous regime."),

    (17, "If it is intermittent, would it be frequent rather than rare?",
     "The pressure being managed is annual and predictable - one concentrated "
     "visitor season per year at each site - which is what makes it tractable "
     "by a management method rather than requiring engineering. The project's "
     "own interventions are rare events: one installation period per site, "
     "then observation."),

    (18, "Would the impact be irreversible?",
     "No. This is the question the design was built to answer. Every component "
     "is demountable and removable without lasting trace; nothing is "
     "excavated beyond post footings; nothing is extracted, dredged, cleared "
     "or imported. The project also carries an explicit stopping rule: if "
     "first-season monitoring shows an installation is causing harm, it is "
     "altered or removed - which is possible only because it was built to be."),

    (19, "Would it be difficult to avoid, reduce, repair or compensate the "
         "effect?",
     "No, and that is a consequence of modularity rather than good fortune. A "
     "single boardwalk section can be repaired or realigned without touching "
     "the rest, a marker line can be moved, and the whole installation can be "
     "removed. The two-stage design sign-off - activity lead for technical "
     "content, authority owner for the site - means changes are agreed by the "
     "body that will live with them."),

    (0, "Other comments/information",
     "AMMOS is an environmental project whose purpose is to reduce erosion and "
     "trampling damage on Mediterranean sediment shores under tourism "
     "pressure. Every site is characterised before intervention and monitored "
     "through a full visitor season after, so the project evidences its own "
     "effect rather than asserting it. Checklist C is not provided because "
     "neither threshold is met: nothing in Infrastructures and works, and the "
     "installations are demountable, not five-year investments."),
]


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Cm(1.6)
        n = self.d.styles["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(10.5)

    def h(self, t, size=13, colour=NAVY, space=12):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(space)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(clean(t))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), colour

    def p(self, t, size=10.5, colour=None, italic=False, bold=False):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(5)
        r = par.add_run(clean(t))
        r.font.size, r.italic, r.bold = Pt(size), italic, bold
        if colour:
            r.font.color.rgb = colour

    def table(self, head, rows, widths=None, small=8.5, flag=None):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, x in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(clean(x))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(small), NAVY
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(clean(v))
                r.font.size = Pt(small)
                if flag is not None and i == flag:
                    r.bold = True
                    r.font.color.rgb = {"YES": GREEN, "NO": GREY,
                                        "N/A": AMBER}.get(str(v).strip(), GREY)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t


def main():
    fails = []
    for label, ans, text in S2:
        n = len(clean(text))
        if n > LIM:
            fails.append(f"S2 {label} is {n} chars, limit {LIM} - cut "
                         f"{n - LIM}")
        if ans not in ("YES", "NO", "N/A"):
            fails.append(f"S2 {label} answer {ans!r} is not YES/NO/N/A")
    for num, q, text in S3:
        n = len(clean(text))
        if n > LIM:
            fails.append(f"S3 Q{num} is {n} chars, limit {LIM} - cut "
                         f"{n - LIM}")
    # The answers must not contradict text already in the application.
    s2 = {l: a for l, a, _ in S2}
    if s2["Coastal areas"] != "YES":
        fails.append("two of three sites are coastal; a NO here contradicts "
                     "3.1.1")
    if s2["Buffer zone of natural protected area"] != "YES":
        fails.append("3.1.1 already states Gomati sits outside GR4110006 "
                     "which covers almost the whole island perimeter - a NO "
                     "on buffer zone contradicts the form two tabs away")
    # Q18 irreversibility is the load-bearing answer.
    q18 = [t for n, q, t in S3 if n == 18][0]
    for must in ("demountable", "removable", "stopping rule"):
        if must not in q18.lower():
            fails.append(f"S3 Q18 does not mention {must!r} - it is the "
                         f"question the whole design answers")
    blob = " ".join(t for _, _, t in S2) + " ".join(t for _, _, t in S3)
    for forbidden in (r"\bRAS\b", r"Regione Autonoma", r"Hakanson", r"xxxx",
                      r"\[[A-Z][A-Z0-9_ ]{2,}\]"):
        if re.search(forbidden, blob, re.I):
            fails.append(f"text contains {forbidden!r}")
    assert not fails, "NOT SHIPPING:\n  - " + "\n  - ".join(fails)

    D = Doc()
    D.h("AMMOS — Checklist B, Sections 2 and 3", size=15)
    D.p(f"Generated {dt.date.today():%d/%m/%Y}. Section 2 is eleven location "
        f"categories, each YES / NO / N/A with a 500-character Specify box. "
        f"Section 3 is nineteen significance questions plus Other comments, "
        f"all 500. Thirty-one boxes, every one checked against its counter.",
        size=9.5, colour=GREY, italic=True)

    D.h("The posture, and why four answers are YES", size=13)
    D.p("Answer YES wherever a site genuinely touches a category, then bound "
        "it precisely in the Specify box. Our own application already says "
        "so: section 3.1.1, which is in the form, states that Gomati \"sits "
        "outside Natura 2000 site GR4110006, which covers almost the whole "
        "coastal and marine perimeter of the island\". Ticking NO against "
        "buffer zone would contradict text an assessor can read two tabs "
        "away, and internal inconsistency is worse than any individual YES.",
        size=10.5)
    D.p("It is also our own argument. The whole premise of AMMOS is that a "
        "habitat under pressure which falls between two designations is "
        "exactly the case a transferable method is for. Denying the proximity "
        "throws the premise away.", size=10.5)
    D.p("So coastal, marine, buffer zone and vulnerable landscape are YES. "
        "Densely populated, cultural heritage, natural protected area, "
        "wetlands, estuaries and special biodiversity area are NO because "
        "they are false, not because NO is comfortable. Mountain areas is "
        "N/A.", size=10.5, bold=True)

    D.h("Section 2 — Location", size=14)
    D.table(["Category", "Answer", "Chars"],
            [[l, a, f"{len(clean(t))}/{LIM}"] for l, a, t in S2],
            widths=[8.0, 2.0, 2.0], flag=1)
    for label, ans, text in S2:
        D.h(f"{label}  —  {ans}", size=11, space=10)
        D.p(f"Specify ({len(clean(text))}/{LIM}):", size=9, colour=GREY,
            italic=True)
        D.p(text)

    D.h("Section 3 — significance of the effects", size=14)
    D.p("Nineteen questions that all ask, in different words, how big and how "
        "permanent the effect is. The answers run one way throughout: small, "
        "local, temporary, intermittent, reversible, avoidable. That "
        "consistency is the point - an assessor reading nineteen boxes is "
        "looking for the one that does not fit.", size=10.5)
    for num, q, text in S3:
        D.h(f"{'Other comments' if not num else f'Question {num}'}", size=11,
            space=10)
        D.p(q, size=10, italic=True, colour=GREY)
        D.p(f"{len(clean(text))}/{LIM}:", size=9, colour=GREY, italic=True)
        D.p(text)

    D.p("")
    D.h("One wording note", size=12)
    D.p("Question 11 in your export reads \"What is the odd that the foreseen "
        "effect(s) will actually take place?\" - I have read that as "
        "likelihood, which is what it must mean. The answer distinguishes the "
        "two cases honestly: the positive effect is likely but measured "
        "rather than assumed, and the negative effects are certain but "
        "trivial. Saying \"certain\" about a small effect reads better than "
        "saying \"unlikely\" about anything.", size=10.5)

    D.d.save(OUT)
    print(f"wrote {Path(OUT).name}")
    print("\nSection 2:")
    for l, a, t in S2:
        print(f"   {a:4s} {len(clean(t)):>4}/{LIM}  {l}")
    print("\nSection 3:")
    for n, q, t in S3:
        print(f"   {('extra' if not n else f'Q{n}'):6s} "
              f"{len(clean(t)):>4}/{LIM}")
    print(f"\n   {len(S2)} + {len(S3)} = {len(S2) + len(S3)} boxes, all within "
          f"{LIM}")


if __name__ == "__main__":
    main()
