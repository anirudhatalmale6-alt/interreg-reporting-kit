#!/usr/bin/env python3
"""Procurement tracker v5: add contract 2 Part B, and move the validation
one column right.

    python3 build_tracker_v5.py

TWO CHANGES, BOTH EVIDENCED.

1. SIX ROWS THAT WERE NEVER ITEMISED.

   Six Lead Beneficiary deliverables carry procurement budget in the approved
   technical sheet and have no row in the tracker: 3.3, 3.4, 4.1, 4.3, 4.4,
   5.4. Together 139.000,00 gross.

   It reconciles exactly, which is why I am confident rather than guessing.
   Contract 2's tender total in the tracker is 257.327,42. The rows already
   present for contract 2 sum to 145.230,65 - which is what their own
   description calls "PART A". So Part B is 112.096,77, and 139.000 / 1,24 is
   112.096,77 to the cent. Part B is inside the contract total and was never
   broken down per deliverable.

   The six new rows copy contract 2's procedure, tender total, planned
   publication and pre-approval answer from the Part A rows, because it is the
   same tender. Planned signature is copied too. Nothing is invented: every
   field is either computed from the approved budget or copied from a sibling
   row.

2. THE DATA VALIDATION IS ONE COLUMN LEFT OF WHERE IT BELONGS.

   The READ ME promises "the cells will warn you if you type something that is
   not a date". They warn the wrong columns.

       date      on K, L, M, S     should be L, M, N, T
       list Y,N  on O, P           should be P, Q
       decimal   on N              should be O

   Every rule is shifted by exactly one, which makes it one mistake rather
   than three. Three independent things say so and they all agree with each
   other and disagree with the validation:

       the HEADERS            L/M/N are the three actual dates, O is the
                              contract amount, P and Q are the two Y/N answers
       the NUMBER FORMATS     L, M, N and T are dd/mm/yyyy; O is #,##0.00
       the STATUS FORMULA     ISNUMBER(N) for signature, ISNUMBER(O) for
                              amount, UPPER(Q)="Y" for blocked

   And the live data agrees: twenty rows already hold a DATE in N, which has a
   decimal rule on it, and a NUMBER in O, which has a "Y,N" list rule on it.

   This is not cosmetic. STATUS decides COMPLETED on ISNUMBER(N). A date typed
   as text is not a number, so a signed contract would sit at NOT STARTED for
   ever and nobody would see it. The date rule on N is exactly the guard
   against that, and it was pointing at the wrong cell.

WHAT IS DELIBERATELY NOT CHANGED

The VAT treatment. Twenty-one Greek rows are the budget divided by exactly
1,24 and twenty-four Albanian rows equal it exactly. That is a rule - Greek
eligible cost includes non-recoverable VAT, Albanian IPA contracts are exempt -
not forty-five errors, and it is the client's call to confirm, not mine to
"fix". The six new rows follow the same rule because they are Greek.

HOW IT IS WRITTEN

openpyxl loads this workbook with no extension warnings, so the conditional
formatting and the validation round-trip intact. The risks are the ones
openpyxl does not handle for you, and each is dealt with explicitly:

  inserting rows does NOT rewrite formulas, so U, V and W are regenerated for
  every data row from the row-4 template. Safe because all three reference
  only cells in their own row.

  inserting rows does NOT extend conditional formatting, validation or the
  autofilter. All three are re-stated over the new last row.

  inserting rows does NOT carry style or the locked/unlocked flag. Both are
  copied cell by cell from a Part A sibling, so the pale-yellow editable
  columns stay editable and the rest stay locked.
"""
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

SRC = Path("/var/lib/freelancer/projects/40536008/"
           "Procurement_Tracker_EPIRUS_MEDEYE_v4.xlsx")
DST = Path("/var/lib/freelancer/projects/40536008/"
           "Procurement_Tracker_EPIRUS_MEDEYE_v5.xlsx")
SHEET = "EPIRUS-MEDEYE"
README = "READ ME FIRST"

FIRST_DATA = 4
VAT = 1.24
CONTRACT = 2
TEMPLATE_ROW = 12          # the last Part A row - same contract, same tender
INSERT_AFTER = 12

# deliverable -> (title as the sheet writes it, gross budget from the
# approved technical sheet version 1.1, export 96606)
PART_B = [
    ("3.3", "Development and Customization of Telemedicine Software", 20_000),
    ("3.4", "Integration of Telemedicine Systems with Health Record", 25_000),
    ("4.1", "Design and Development of Training Programs for Teleme", 25_000),
    ("4.3", "Creation of Educational Materials and E-learning Modul", 24_000),
    ("4.4", "Establishment of a Technical Support and Helpline", 30_000),
    ("5.4", "Analysis and Optimization of Telemedicine Services Bas", 15_000),
]
PART_A_SUM = 145_230.65
TENDER_TOTAL = 257_327.42

# column -> what the header, the number format and the STATUS formula all
# agree it holds. The validation currently sits one column to the left.
DV_FIX = {
    "date": ["L", "M", "N", "T"],
    "list": ["P", "Q"],
    "decimal": ["O"],
}


def eur(x):
    return f"{x:,.2f}".translate(str.maketrans({",": ".", ".": ","}))


def main():
    assert SRC.exists(), f"missing {SRC}"
    wb = openpyxl.load_workbook(SRC)
    ws = wb[SHEET]
    last = ws.max_row

    # ---- prove the arithmetic against the workbook before touching it -----
    c2 = [r for r in range(FIRST_DATA, last + 1)
          if str(ws.cell(row=r, column=1).value or "").startswith("LB")
          and str(ws.cell(row=r, column=2).value).strip() in ("2", "2.0")]
    part_a = sum(ws.cell(row=r, column=7).value or 0 for r in c2)
    total = next(ws.cell(row=r, column=8).value for r in c2)
    gross = sum(b for _, _, b in PART_B)
    part_b = round(total - part_a, 2)

    fails = []
    if abs(part_a - PART_A_SUM) > 0.02:
        fails.append(f"Part A reads {eur(part_a)}, expected {eur(PART_A_SUM)}")
    if abs(total - TENDER_TOTAL) > 0.02:
        fails.append(f"tender total reads {eur(total)}")
    if abs(part_b - gross / VAT) > 0.02:
        fails.append(
            f"Part B is {eur(part_b)} but the six missing deliverables come "
            f"to {eur(gross / VAT)} ex VAT - they do not reconcile, so the "
            f"premise of this whole script is wrong")
    have = {str(ws.cell(row=r, column=5).value or "").split("|")[0].strip()
            for r in range(FIRST_DATA, last + 1)
            if str(ws.cell(row=r, column=1).value or "").startswith("LB")}
    for d, _, _ in PART_B:
        if d in have:
            fails.append(f"deliverable {d} already has a Lead Beneficiary "
                         f"row - adding it again would double count")
    if ws.max_row != INSERT_AFTER + len(
            [r for r in range(INSERT_AFTER + 1, last + 1)]) :
        pass  # row count is informational only
    assert not fails, "NOT WRITING:\n  - " + "\n  - ".join(fails)

    print(f"Part A {eur(part_a)} + Part B {eur(part_b)} = {eur(total)}")
    print(f"six deliverables gross {eur(gross)} / 1,24 = {eur(gross / VAT)}")

    # ---- insert the six rows, keeping the contract-2 block together -------
    ws.insert_rows(INSERT_AFTER + 1, amount=len(PART_B))
    tmpl = TEMPLATE_ROW
    for i, (deliv, title, budget) in enumerate(PART_B):
        r = INSERT_AFTER + 1 + i
        for ci in range(1, ws.max_column + 1):
            src = ws.cell(row=tmpl, column=ci)
            dst = ws.cell(row=r, column=ci)
            dst._style = copy(src._style)
            dst.protection = copy(src.protection)
        ws.cell(row=r, column=1).value = ws.cell(row=tmpl, column=1).value
        ws.cell(row=r, column=2).value = ws.cell(row=tmpl, column=2).value
        ws.cell(row=r, column=3).value = f"PART B: {eur(part_b)}"
        ws.cell(row=r, column=4).value = ws.cell(row=tmpl, column=4).value
        ws.cell(row=r, column=5).value = f"{deliv} | {title}"
        ws.cell(row=r, column=6).value = ws.cell(row=tmpl, column=6).value
        ws.cell(row=r, column=7).value = round(budget / VAT, 2)
        ws.cell(row=r, column=8).value = total
        ws.cell(row=r, column=9).value = ws.cell(row=tmpl, column=9).value
        ws.cell(row=r, column=10).value = ws.cell(row=tmpl, column=10).value
        ws.cell(row=r, column=11).value = ws.cell(row=tmpl, column=11).value
        for ci in range(12, 21):            # L..T stay for the partner
            ws.cell(row=r, column=ci).value = None
        ws.cell(row=r, column=16).value = ws.cell(row=tmpl, column=16).value

    new_last = ws.max_row

    # ---- regenerate the three calculated columns for EVERY data row ------
    # openpyxl does not rewrite formulas when rows move, and all three refer
    # only to their own row, so rebuilding them from the template is exact.
    # A plain str.replace of "4" would also rewrite the 1.24 inside the W
    # formula - 1.213, 1.220, 1.256 - silently changing the VAT rate per row.
    # It did, on the first run. Only a CELL REFERENCE may be renumbered: a
    # column letter immediately before the row number, with no digit after.
    REF = re.compile(rf"(\$?[A-Z]{{1,3}}\$?){FIRST_DATA}(?![0-9])")
    tpl = {c: ws.cell(row=FIRST_DATA, column=i).value
           for c, i in (("U", 21), ("V", 22), ("W", 23))}
    for col, i in (("U", 21), ("V", 22), ("W", 23)):
        base = tpl[col]
        assert isinstance(base, str) and base.startswith("="), col
        for r in range(FIRST_DATA, new_last + 1):
            ws.cell(row=r, column=i).value = REF.sub(rf"\g<1>{r}", base)

    # Every constant in the templates must survive untouched in every row.
    CONSTANTS = {"U": [], "V": [], "W": ["1.24", "#,##0"]}
    broken = []
    for col, i in (("U", 21), ("V", 22), ("W", 23)):
        for r in range(FIRST_DATA, new_last + 1):
            f = ws.cell(row=r, column=i).value
            for k in CONSTANTS[col]:
                if k not in f:
                    broken.append(f"{col}{r} lost the constant {k!r}")
            if f.count(f"{r}") < 2:
                broken.append(f"{col}{r} was not renumbered")
    assert not broken, ("FORMULA REWRITE DAMAGED THE SHEET:\n  - "
                        + "\n  - ".join(broken[:10]))

    # ---- extend conditional formatting over the new rows -----------------
    from openpyxl.formatting.formatting import ConditionalFormattingList
    old_cf = ws.conditional_formatting
    new_cf = ConditionalFormattingList()
    for cf in old_cf:
        sq = str(cf.sqref).replace(str(last), str(new_last))
        for rule in cf.rules:
            new_cf.add(sq, rule)
    ws.conditional_formatting = new_cf

    # ---- restate the validation, one column to the RIGHT -----------------
    old = {dv.type: dv for dv in ws.data_validations.dataValidation}
    ws.data_validations.dataValidation = []
    for kind, cols in DV_FIX.items():
        src = old.get(kind)
        dv = DataValidation(
            type=src.type, operator=src.operator,
            formula1=src.formula1, formula2=src.formula2,
            allow_blank=True, showErrorMessage=True, showInputMessage=True,
            errorTitle=src.errorTitle, error=src.error,
            promptTitle=src.promptTitle, prompt=src.prompt)
        for c in cols:
            dv.add(f"{c}{FIRST_DATA}:{c}{new_last}")
        ws.add_data_validation(dv)

    ws.auto_filter.ref = f"A3:V{new_last}"

    # ---- note it in the READ ME so nobody has to ask what changed --------
    rm = wb[README]
    row = rm.max_row + 2
    rm.cell(row=row, column=2).value = "What changed in v5 - 11 October 2026"
    rm.cell(row=row + 1, column=2).value = (
        f"Added six rows for contract 2, Part B - deliverables "
        f"{', '.join(d for d, _, _ in PART_B)}. These carry procurement "
        f"budget in the approved technical sheet (version 1.1, export 96606) "
        f"and had no row here. {eur(gross)} gross, {eur(part_b)} ex VAT, "
        f"which is exactly the difference between contract 2's tender total "
        f"of {eur(total)} and the Part A rows at {eur(part_a)}.")
    rm.cell(row=row + 2, column=2).value = (
        "Moved the data validation one column right. The date rule now sits "
        "on ACTUAL publication, ACTUAL offers deadline, ACTUAL signature and "
        "Updated on; the Y/N rule on Pre-approval obtained and Blocked?; the "
        "number rule on Contract amount incl VAT. It was one column left of "
        "all of those, which meant the signature date - the field STATUS uses "
        "to decide COMPLETED - had no date check on it at all.")
    rm.cell(row=row + 3, column=2).value = (
        "Not changed: the VAT treatment. Greek rows are the budget divided by "
        "1,24 and Albanian rows equal it, which is a rule rather than an "
        "error and needs the FLC to confirm, not a script to overwrite. The "
        "six new rows are Greek and follow it.")
    rm.cell(row=row + 4, column=2).value = (
        f"Still open: {eur(258_779.04 - total)} between the published tender "
        f"of {eur(258_779.04)} ex VAT and Part A plus Part B. One deliverable "
        f"share that has never been allocated.")

    wb.save(DST)
    print(f"wrote {DST}")
    print(f"  rows {last} -> {new_last}")
    return DST, new_last


if __name__ == "__main__":
    main()
