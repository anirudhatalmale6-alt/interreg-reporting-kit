"""
EPIRUSMEDEYE - claims ledger: check every uploaded expense table, and show how
much money is LEFT per partner and deliverable.

Setup (PyCharm terminal, once):
    pip install openpyxl

Use:
    1. Put every "Πίνακας Δαπανών ... .xlsx" you download into one folder
       (sub-folders are fine - it searches them all). Keep old claims there
       too: the "left to spend" figures add up ALL claims in the folder.
    2. Check the four paths in CONFIG below.
    3. Run. Open the file named in OUTPUT_FILE.

What it reads:
    TECH_SHEET  the approved technical sheet export (version 1.2,
                "2-interreg_report129557.xlsx")
    TRACKER     the procurement tracker (v5). Optional - set to None to skip
                the contract checks.
    CLAIMS_DIR  the folder with the expense tables.

What it checks, per expense line:
    - the partner and deliverable exist in the approved budget
    - the running total per partner x deliverable x cost category never goes
      above the approved amount
    - the invoice is dated on/after the project start
    - a split invoice adds up exactly to the invoice total
    - procured costs (external services, equipment) have a SIGNED contract in
      the tracker for that deliverable, signed on or before the invoice date
    - the flat rates claimed are the same percentages as the approved budget

It refuses to mix projects: any expense table whose MIS code is not
PROJECT_MIS is skipped and listed on the "Files" sheet. The same claim (same
"id δελτίου") found twice is counted once.

Colours: green = fine, amber = look at it, red = a problem to raise before
approval.
"""

import datetime as dt
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================== CONFIG ==============================
TECH_SHEET = r"D:\OFFICE\EPIRUSMEDEYE\2-interreg_report129557.xlsx"
TRACKER = r"D:\OFFICE\EPIRUSMEDEYE\Procurement_Tracker_EPIRUS_MEDEYE_v5.xlsx"
CLAIMS_DIR = r"D:\OFFICE\EPIRUSMEDEYE\CLAIMS"
OUTPUT_FILE = r"D:\OFFICE\EPIRUSMEDEYE\claims_ledger.xlsx"

PROJECT_MIS = "6010984"
PROJECT_START = dt.date(2025, 10, 13)
WARN_AT = 0.90            # amber when a line is 90% or more spent
# ====================================================================

GREEN, AMBER, RED = "C6EFCE", "FFEB9C", "FFC7CE"
LABEL = {GREEN: "OK", AMBER: "CHECK", RED: "PROBLEM"}

# claim category code -> budget column in the technical sheet
CATEGORY = {
    "Α.25": "staff", "Β.24": "staff",          # staff (real cost / flat)
    "Β.26": "overheads",                         # office & admin, flat
    "Β.30": "travel",                            # travel & stay, flat
    "Α.33": "services",                          # external expertise
    "Α.29": "equipment",
}
FLAT = {"Β.24", "Β.26", "Β.30"}
PROCURED = {"services", "equipment"}
COLS = ["staff", "overheads", "travel", "services", "equipment",
        "infrastructure"]


def norm(name):
    """Compare organisation names regardless of quotes, spaces and case."""
    return re.sub(r"[\s'\"΄`’‘.,()-]+", "", str(name)).upper()


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.datetime.strptime(str(v).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def code_of(cat):
    m = re.search(r"([ΑΒ]\.\d+)", str(cat))
    return m.group(1) if m else str(cat).strip()


# ------------------------------------------------------------ the budget
def read_budget(path):
    """Partner -> deliverable -> column amounts, from the tech sheet export.

    Self-checking: every partner's deliverables must add up to the
    'ΓΕΝΙΚΟ ΣΥΝΟΛΟ' printed in the sheet, or it stops. A misread budget is
    worse than none, because everything else is measured against it.
    """
    ws = load_workbook(path, data_only=True).active
    partners, budget, printed = [], {}, {}
    cur = None
    for row in ws.iter_rows(values_only=True):
        c = ["" if x is None else x for x in row]
        t = str(c[0])
        head = re.match(r"\s*(?:Κύριος Δικαιούχος|P\d)\s*\d+\s*\((.+?)\)\s*"
                        r"\(", t)
        if head:
            cur = head.group(1).strip()
            partners.append(cur)
            budget[cur] = {}
            continue
        if t.startswith("ΓΕΝΙΚΟ ΣΥΝΟΛΟ") and cur:
            printed[cur] = [x for x in c if isinstance(x, (int, float))][-1]
            continue
        md = re.match(r"\s*Παραδοτέο\s*(\d+)\.\s*(\d+)", t)
        if md and cur:
            n = [x for x in c[2:] if isinstance(x, (int, float))]
            budget[cur][f"{md.group(1)}.{md.group(2)}"] = dict(
                zip(COLS + ["total"], [round(v, 2) for v in n[:7]]))
    assert partners, f"no partner blocks found in {path}"
    for p in partners:
        s = sum(d["total"] for d in budget[p].values())
        assert abs(s - printed[p]) < 0.05, (
            f"budget misread for {p}: deliverables add to {s:,.2f}, sheet "
            f"prints {printed[p]:,.2f}")
    return partners, budget


# ----------------------------------------------------------- the tracker
def read_tracker(path, partners):
    """Contracts per partner, keyed by the order PB1..PB4 = budget order."""
    if not path:
        return None
    ws = load_workbook(path, data_only=True)["EPIRUS-MEDEYE"]
    out = defaultdict(list)
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r[0]:
            continue
        m = re.search(r"PB(\d)", str(r[0]))
        if not m or int(m.group(1)) > len(partners):
            continue
        p = partners[int(m.group(1)) - 1]
        out[p].append(dict(
            no=str(r[1]).strip(), buys=str(r[2] or "")[:80],
            deliv=str(r[4]).split("|")[0].strip(),
            this_ex_vat=r[6] or 0, tender_ex_vat=r[7] or 0,
            signed=as_date(r[13]) if r[13] else None,
            amount=float(r[14]) if isinstance(r[14], (int, float)) else None))
    return out


# ------------------------------------------------------------- the claims
def read_claim(path):
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or "ΠΙΝΑΚΑΣ ΔΑΠΑΝΩΝ" not in " ".join(
            str(x) for r in rows[:4] for x in r if x):
        return None
    meta = {str(r[0]).strip(): r[5] for r in rows[:12] if r and r[0]}
    lines, flats = [], []
    for r in rows[16:]:
        if r and r[0] and str(r[0]).startswith("ΣΥΝΟΛΙΚΑ"):
            break
        if not r or len(r) < 15:
            continue
        cat = code_of(r[11]) if r[11] else ""
        if r[0] not in (None, ""):
            lines.append(dict(
                inv=str(r[0]).strip(), inv_date=as_date(r[1]),
                supplier=str(r[2] or "").strip(),
                inv_total=float(r[4] or 0), pay_date=as_date(r[7]),
                alloc=str(r[10] or "").strip(), cat=cat,
                deliv=str(r[12] or "").strip(),
                declared=float(r[14] or 0)))
        elif cat:
            flats.append(dict(cat=cat, declared=float(r[14] or 0)))
    return dict(path=path, mis=str(meta.get("Κωδικός ΟΠΣ", "")).strip(),
                claim=str(meta.get("id δελτίου", "")).strip(),
                partner=str(meta.get("Δικαιούχος", "")).strip(),
                lines=lines, flats=flats)


# ---------------------------------------------------------------- checks
def run():
    partners, budget = read_budget(TECH_SHEET)
    by_norm = {norm(p): p for p in partners}
    contracts = read_tracker(TRACKER, partners)

    files, claims, seen = [], [], set()
    for f in sorted(Path(CLAIMS_DIR).rglob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        try:
            c = read_claim(f)
        except Exception as e:                     # unreadable: say so
            files.append((str(f), "", "", "NOT READ", f"error: {e}"))
            continue
        if c is None:
            continue
        if not c["lines"] and not c["flats"]:
            files.append((str(f), c["mis"], c["claim"], "SKIPPED",
                          "empty expense table"))
            continue
        if c["mis"] != PROJECT_MIS:
            files.append((str(f), c["mis"], c["claim"], "SKIPPED",
                          f"different project (MIS {c['mis']})"))
            continue
        if c["claim"] in seen:
            files.append((str(f), c["mis"], c["claim"], "SKIPPED",
                          "same claim already read - counted once"))
            continue
        p = by_norm.get(norm(c["partner"]))
        if not p:
            files.append((str(f), c["mis"], c["claim"], "NOT MATCHED",
                          f"partner '{c['partner']}' not in the budget"))
            continue
        seen.add(c["claim"])
        c["budget_partner"] = p
        claims.append(c)
        files.append((str(f), c["mis"], c["claim"], "READ",
                      f"{p}: {len(c['lines'])} expense lines"))

    used = defaultdict(float)          # (partner, deliv, column) -> claimed
    flat_used = defaultdict(float)     # (partner, column) -> claimed
    expenses, rate_lines = [], []

    for c in claims:
        p = c["budget_partner"]
        parts = defaultdict(list)
        for x in c["lines"]:
            parts[(x["inv"], x["supplier"], x["inv_total"])].append(x)
        for x in c["lines"]:
            why, flag = [], GREEN
            col = CATEGORY.get(x["cat"])
            line = budget[p].get(x["deliv"])
            if col is None:
                flag, _ = RED, why.append(f"unknown category {x['cat']}")
            elif line is None:
                flag, _ = RED, why.append(f"deliverable {x['deliv']} is not "
                                          f"in {p}'s approved budget")
            else:
                used[(p, x["deliv"], col)] += x["declared"]
                cap, u = line[col], used[(p, x["deliv"], col)]
                if u > cap + 0.005:
                    flag = RED
                    why.append(f"{x['deliv']} {col}: {u:,.2f} claimed so far "
                               f"against {cap:,.2f} approved")
                elif cap and u >= WARN_AT * cap:
                    flag = AMBER if flag == GREEN else flag
                    why.append(f"{x['deliv']} {col} now {100*u/cap:.0f}% "
                               f"spent ({cap-u:,.2f} left)")
            if x["inv_date"] and x["inv_date"] < PROJECT_START:
                flag = RED
                why.append("invoice dated before the project start")
            grp = parts[(x["inv"], x["supplier"], x["inv_total"])]
            if len(grp) > 1:
                s = sum(g["declared"] for g in grp)
                if abs(s - x["inv_total"]) > 0.01:
                    flag = RED
                    why.append(f"split invoice: parts {s:,.2f} vs invoice "
                               f"{x['inv_total']:,.2f}")
            if contracts is not None and col in PROCURED:
                cs = [k for k in contracts.get(p, []) if k["deliv"] == x["deliv"]]
                signed = [k for k in cs if k["signed"]]
                if not cs:
                    flag = AMBER if flag == GREEN else flag
                    why.append("no contract in the tracker for this "
                               "deliverable")
                elif not signed:
                    flag = RED
                    why.append("contract in the tracker is not signed yet")
                elif x["inv_date"] and x["inv_date"] < min(
                        k["signed"] for k in signed):
                    flag = RED
                    why.append("invoice dated before the contract signature")
                else:
                    hit = [k for k in signed if k["amount"] and
                           abs(k["amount"] - x["inv_total"]) < 0.01]
                    why.append(f"contract {hit[0]['no']} (exact amount)" if hit
                               else "signed contract exists for this "
                                    "deliverable")
            expenses.append([LABEL[flag], c["claim"], p, x["inv"],
                             x["inv_date"], x["supplier"][:60], x["inv_total"],
                             x["pay_date"], x["cat"], x["deliv"],
                             x["declared"], "; ".join(why) or "fine", flag])
        # flat rates must be the approved percentages
        direct = sum(x["declared"] for x in c["lines"])
        staff = sum(x["declared"] for x in c["lines"]
                    if CATEGORY.get(x["cat"]) == "staff")
        a = {k: sum(d[k] for d in budget[p].values()) for k in COLS}
        a_direct = a["services"] + a["equipment"] + a["infrastructure"]
        for f in c["flats"]:
            col = CATEGORY.get(f["cat"])
            flat_used[(p, col)] += f["declared"]
            if f["cat"] == "Β.24":
                staff += f["declared"]
                got, want = f["declared"] / direct if direct else 0, (
                    a["staff"] / a_direct if a_direct else 0)
                base = "direct costs"
            else:
                st = staff or sum(g["declared"] for g in c["flats"]
                                  if g["cat"] == "Β.24")
                got = f["declared"] / st if st else 0
                want = a[col] / a["staff"] if a["staff"] else 0
                base = "staff"
            ok = abs(got - want) < 0.0001
            rate_lines.append([LABEL[GREEN if ok else RED], c["claim"], p,
                               f["cat"], f["declared"], f"{100*got:.3f}%",
                               f"{100*want:.3f}%", f"% of {base}",
                               GREEN if ok else RED])

    # ---------------------------------------------------- write the workbook
    wb = Workbook()
    s = wb.active
    s.title = "Partners"
    s.append(["Partner", "Approved", "Claimed to date", "Left", "% spent"])
    for p in partners:
        appr = sum(d["total"] for d in budget[p].values())
        cl = sum(v for (pp, _, _), v in used.items() if pp == p) + sum(
            v for (pp, _), v in flat_used.items() if pp == p)
        s.append([p, appr, round(cl, 2), round(appr - cl, 2),
                  round(100 * cl / appr, 1) if appr else None])

    left = wb.create_sheet("Left per deliverable")
    left.append(["Partner", "Deliverable", "Cost category", "Approved",
                 "Claimed to date", "Left", "% spent", ""])
    for p in partners:
        for dl, d in sorted(budget[p].items(),
                            key=lambda kv: [int(n) for n in kv[0].split(".")]):
            for col in ("services", "equipment", "infrastructure"):
                if d[col] or used.get((p, dl, col)):
                    u = round(used.get((p, dl, col), 0), 2)
                    pct = 100 * u / d[col] if d[col] else 100
                    flag = RED if u > d[col] + 0.005 else (
                        AMBER if pct >= 100 * WARN_AT else GREEN)
                    left.append([p, dl, col, d[col], u, round(d[col] - u, 2),
                                 round(pct, 1), flag])
        a = {k: sum(d[k] for d in budget[p].values()) for k in COLS}
        for col in ("staff", "overheads", "travel"):
            if a[col] or flat_used.get((p, col)):
                u = round(flat_used.get((p, col), 0) + sum(
                    v for (pp, _, cc), v in used.items()
                    if pp == p and cc == col), 2)
                pct = 100 * u / a[col] if a[col] else 100
                flag = RED if u > a[col] + 0.005 else (
                    AMBER if pct >= 100 * WARN_AT else GREEN)
                left.append([p, "(whole partner)", col + " (flat rate)",
                             round(a[col], 2), u, round(a[col] - u, 2),
                             round(pct, 1), flag])

    e = wb.create_sheet("Expenses")
    e.append(["Flag", "Claim id", "Partner", "Invoice", "Invoice date",
              "Supplier", "Invoice total", "Paid on", "Category",
              "Deliverable", "Declared", "Why", ""])
    for r in expenses:
        e.append(r)

    rt = wb.create_sheet("Flat rates")
    rt.append(["Flag", "Claim id", "Partner", "Category", "Claimed",
               "Rate claimed", "Rate approved", "Base", ""])
    for r in rate_lines:
        rt.append(r)

    if contracts is not None:
        ct = wb.create_sheet("Contracts")
        ct.append(["Partner", "Contract", "Buys", "Deliverable",
                   "This deliverable ex VAT", "Tender total ex VAT",
                   "Signed", "Contract amount incl VAT"])
        for p in partners:
            for k in contracts.get(p, []):
                ct.append([p, k["no"], k["buys"], k["deliv"],
                           round(k["this_ex_vat"], 2), k["tender_ex_vat"],
                           k["signed"], k["amount"]])

    fl = wb.create_sheet("Files")
    fl.append(["File", "MIS", "Claim id", "Result", "Note"])
    for r in files:
        fl.append(list(r))

    # colour + tidy
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        last = ws.max_column
        flagcol = last if ws[1][last - 1].value in ("", None) else None
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
                if isinstance(c.value, dt.date):
                    c.number_format = "dd/mm/yyyy"
            if flagcol and row[flagcol - 1].value in (GREEN, AMBER, RED):
                fill = PatternFill("solid", fgColor=row[flagcol - 1].value)
                row[0].fill = fill
                if ws.title == "Left per deliverable":
                    for c in row[3:7]:
                        c.fill = fill
                row[flagcol - 1].value = None
        for i in range(1, last + 1):
            width = max((len(str(c.value)) for c in ws[get_column_letter(i)]
                         if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2,
                                                                   60)
        if ws.title == "Expenses":
            ws.column_dimensions["L"].width = 90
            for row in ws.iter_rows(min_row=2):
                row[11].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUTPUT_FILE)

    n = defaultdict(int)
    for r in expenses:
        n[r[0]] += 1
    print(f"claims read: {len(claims)}  |  expense lines: {len(expenses)}  "
          f"(OK {n['OK']}, CHECK {n['CHECK']}, PROBLEM {n['PROBLEM']})")
    bad_rates = [r for r in rate_lines if r[0] == "PROBLEM"]
    print(f"flat rates checked: {len(rate_lines)}  |  wrong: {len(bad_rates)}")
    for r in files:
        if r[3] != "READ":
            print(f"  {r[3]}: {Path(r[0]).name} - {r[4]}")
    print(f"report: {OUTPUT_FILE}")
    return expenses, used, flat_used, files


if __name__ == "__main__":
    run()
