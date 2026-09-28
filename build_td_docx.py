#!/usr/bin/env python3
"""The approved budget as a Greek report annex, straight from the MIS PDF.

    python3 build_td_docx.py report_140109.pdf

YOUR QUESTION: "how can I handle the exported report after, so I can produce
a report from it? Ask claude?"

No. Ask Claude for the JUDGEMENTS and let a script do the tables.

A budget annex is the worst possible thing to produce by hand or by prompt.
It is forty numbers that all have to agree with each other, in a document a
Ministry may check against the MIS. There is no judgement in it at all - it
is transcription - and transcription is precisely what a model does
imperfectly and a script does perfectly.

So the division of labour on the MINOR-MED report is:

    this script        the identity block, the budget tables, the version
                       history, the reporting calendar. Every figure read
                       from the approved technical sheet and cross-checked
                       39 ways before the document is written.

    build_charts.py    the charts, from the procurement tracker.

    build_narrative.py the factual sentences, plus a prompt block of proven
                       figures for the interpretation.

    you                what it means, what is at risk, what you propose.
                       That part is not automatable and should not be.

WHAT THIS DOCUMENT IS NOT

It is not a spending report. The technical sheet is the APPROVED BUDGET.
Reporting it under a heading that says spending tells the reader the money is
gone. The headings here say εγκεκριμένος προϋπολογισμός throughout, and
there is a note in the reviewer block saying so explicitly, because that is
the single easiest way for this annex to mislead.
"""
import datetime as dt
import pathlib
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

from _shared import HERE, ONLY_EPIRUS, ONLY_MINORMED, clean, guard, newest
from read_td import (CATS, analytical, by_partner, header, indicators, money,
                     text_of, verify, work_packages, wp_by_category)

NAVY = RGBColor(0x1F, 0x38, 0x5E)
GREY = RGBColor(0x55, 0x5F, 0x6B)
RED = RGBColor(0xB4, 0x55, 0x3C)

# Which programme's names must never appear in THIS document. Keyed by MIS so
# it cannot be got wrong by hand: the guard for a 6010952 document is the
# EPIRUS name list, and vice versa.
FORBIDDEN_BY_MIS = {"6010952": ONLY_EPIRUS, "6010984": ONLY_MINORMED}

# The reporting calendar he gave me on 28 September: five reports over twenty
# months from 24/06/2026, the same period for the Joint Secretariat and for
# EDEYPY. Stored as the two facts he stated, and the five periods DERIVED
# from them, so a correction to either one cannot leave a stale schedule
# behind in the document.
REPORTING = {"6010952": dict(first_start=dt.date(2026, 6, 24), months=20,
                             reports=5)}


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return d.replace(year=d.year + y, month=m + 1)


def schedule(mis, project_start, project_end):
    cfg = REPORTING.get(mis)
    if not cfg:
        return [], []
    each = cfg["months"] // cfg["reports"]
    assert cfg["months"] % cfg["reports"] == 0, (
        f"{cfg['months']} months do not divide into {cfg['reports']} equal "
        f"reports. Do not guess the period lengths - ask.")
    rows, notes = [], []
    start = cfg["first_start"]

    # THE FOUR MONTHS BEFORE REPORTING STARTS
    #
    # The project begins 13/02/2026 and the reporting calendar begins
    # 24/06/2026. That leaves four months and eleven days of implementation
    # sitting before the first reporting period opens. A reporting chain with
    # a hole in it is the kind of thing a Managing Authority asks about
    # exactly once, so the gap gets its own row rather than being invisible
    # between the project start and report 1.
    if start > project_start:
        gap = (start - project_start).days
        rows.append(("πριν την 1η", f"{project_start:%d/%m/%Y}",
                     f"{start:%d/%m/%Y}", f"{gap} ημέρες",
                     "ΔΕΝ καλύπτεται από τις πέντε εκθέσεις"))
        notes.append((
            "Διάστημα υλοποίησης πριν την έναρξη του ημερολογίου αναφοράς",
            f"Η πράξη ξεκινά {project_start:%d/%m/%Y} και η πρώτη περίοδος "
            f"αναφοράς {start:%d/%m/%Y}, δηλαδή μεσολαβούν {gap} ημέρες "
            f"({Doc.num(gap / 30.44)} μήνες). Δύο πιθανές αναγνώσεις: είτε η "
            f"1η "
            f"έκθεση καλύπτει αναδρομικά από την έναρξη της πράξης, είτε το "
            f"διάστημα καλύπτεται χωριστά. Να επιβεβαιωθεί, γιατί από αυτό "
            f"εξαρτάται ποια δραστηριότητα αναφέρεται στην 1η έκθεση."))

    for i in range(cfg["reports"]):
        a = add_months(start, each * i)
        b = add_months(start, each * (i + 1))
        flag = ""
        if b > project_end:
            flag = (f"λήγει {(b - project_end).days} ημέρες ΜΕΤΑ τη λήξη "
                    f"του έργου ({project_end:%d/%m/%Y})")
            notes.append(
                ("Η 5η περίοδος αναφοράς λήγει μετά τη λήξη του έργου",
                 f"Πέντε τετράμηνες περίοδοι από {start:%d/%m/%Y} οδηγούν σε "
                 f"λήξη της τελευταίας την {b:%d/%m/%Y}, ενώ το έργο λήγει "
                 f"την {project_end:%d/%m/%Y}. Η τελευταία περίοδος αναφοράς "
                 f"εκτείνεται δηλαδή {(b - project_end).days} ημέρες πέραν "
                 f"της λήξης του έργου. Είτε η τελευταία "
                 f"περίοδος είναι συντομότερη, είτε προβλέπεται περίοδος "
                 f"κλεισίματος. Να επιβεβαιωθεί με την Κοινή Γραμματεία πριν "
                 f"δεσμευτεί ημερομηνία υποβολής."))
        rows.append((f"{i + 1} από {cfg['reports']}", f"{a:%d/%m/%Y}",
                     f"{b:%d/%m/%Y}", f"{(b - a).days} ημέρες", flag))
    return rows, notes


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Cm(2.0)
        n = self.d.styles["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(10)

    def h(self, text, size=13, colour=NAVY, space=10):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(space)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(clean(text))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), colour
        return p

    @staticmethod
    def num(x, dp=1):
        """Greek decimal convention: 4,3 rather than 4.3."""
        return f"{x:.{dp}f}".replace(".", ",")

    def p(self, text, size=10, colour=None, italic=False):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run(clean(text))
        r.font.size, r.italic = Pt(size), italic
        if colour:
            r.font.color.rgb = colour
        return par

    def table(self, headings, rows, widths=None, right_from=1):
        t = self.d.add_table(rows=1, cols=len(headings))
        t.style = "Table Grid"
        for i, htxt in enumerate(headings):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(clean(htxt))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), NAVY
            if i >= right_from:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(clean(val))
                r.font.size = Pt(8.5)
                if str(val).startswith("ΣΥΝΟΛ") or str(row[0]).startswith("ΣΥΝΟΛ"):
                    r.bold = True
                if i >= right_from:
                    cells[i].paragraphs[0].alignment = \
                        WD_ALIGN_PARAGRAPH.RIGHT
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t


def main(pdf=None):
    pdf = pdf or newest("*.pdf")
    t = text_of(pdf)
    hdr = header(t)
    fund_rows, fund_tot = by_partner(t)
    wps, wp_grand = work_packages(t)
    wpcc, wpcc_tot = wp_by_category(t)
    parts = analytical(t)
    inds = indicators(t)

    # THE TABLES ARE VERIFIED BEFORE THE DOCUMENT IS WRITTEN, NOT AFTER.
    #
    # A document that has already been saved is a document that can be sent.
    # So the cross-checks run first and the .docx is never created at all if
    # any of them fail - rather than writing a plausible annex and printing a
    # warning that scrolls off the screen.
    class Sink:
        def __call__(self, *a):
            pass
    fails = verify(Sink(), hdr, fund_rows, fund_tot, wps, wp_grand, wpcc,
                   wpcc_tot, parts)
    assert not fails, ("the technical sheet did not pass its own cross-checks, "
                       "so no document was written:\n  " + "\n  ".join(fails))

    mis = hdr["mis"]
    forbidden = FORBIDDEN_BY_MIS.get(mis)
    assert forbidden, (
        f"MIS {mis} is not in FORBIDDEN_BY_MIS, so this document would be "
        f"written with NO contamination guard at all. Add the MIS and the "
        f"other programme's name list before running it.")

    D = Doc()
    D.h(f"{hdr['acronym']} (MIS {hdr['mis']})", size=16)
    D.p(hdr["title"], size=10.5, italic=True)
    D.h("ΠΑΡΑΡΤΗΜΑ: ΕΓΚΕΚΡΙΜΕΝΟΣ ΠΡΟΫΠΟΛΟΓΙΣΜΟΣ ΚΑΤΑ ΕΤΑΙΡΟ ΚΑΙ ΚΑΤΗΓΟΡΙΑ "
        "ΔΑΠΑΝΗΣ", size=12)
    D.p(f"Πηγή: Τεχνικό Δελτίο Πράξης, έκδοση {hdr['version']} "
        f"({hdr['kind']}), υποβολή {hdr['submitted']}, έγκριση "
        f"{hdr['approved']}. Αρχείο {pdf.name}.", size=9, colour=GREY)
    D.p(f"Διάρκεια πράξης: {hdr['start']} έως {hdr['end']} "
        f"({hdr['months']} μήνες).", size=9, colour=GREY)

    # ---------------------------------------------------------------- totals
    D.h("1. Συνολικός εγκεκριμένος προϋπολογισμός")
    D.table(["Πηγή χρηματοδότησης", "Ποσό (EUR)", "Ποσοστό"],
            [["Κοινοτική συνδρομή", money(fund_tot["eu"]),
              f"{fund_tot['eu'] / fund_tot['total']:.2%}".replace(".", ",")],
             ["Εθνική συμμετοχή", money(fund_tot["nat"]),
              f"{fund_tot['nat'] / fund_tot['total']:.2%}".replace(".", ",")],
             ["ΣΥΝΟΛΟ ΕΠΙΛΕΞΙΜΗΣ ΔΑΠΑΝΗΣ", money(fund_tot["total"]),
              "100,00%"]],
            widths=[8.5, 4.5, 3.5])

    # ------------------------------------------------------------- partners
    D.h("2. Εγκεκριμένος προϋπολογισμός κατά δικαιούχο")
    rows = []
    for p in parts:
        f = fund_rows[p["no"] - 1] if p["no"] <= len(fund_rows) else {}
        rows.append([f"{p['no']}. {p['name']}", f.get("country", ""),
                     money(f.get("eu", 0)), money(f.get("nat", 0)),
                     money(p["printed_total"] or 0),
                     f"{(p['printed_total'] or 0) / fund_tot['total']:.1%}"
                     .replace(".", ",")])
    rows.append(["ΣΥΝΟΛΟ", "", money(fund_tot["eu"]), money(fund_tot["nat"]),
                 money(fund_tot["total"]), "100,0%"])
    D.table(["Δικαιούχος", "Χώρα", "Κοινοτική", "Εθνική", "Σύνολο", "%"],
            rows, widths=[6.0, 1.9, 2.4, 2.2, 2.4, 1.3], right_from=2)

    # ------------------------------------------------------- cost categories
    D.h("3. Εγκεκριμένος προϋπολογισμός κατά κατηγορία δαπάνης")
    rows = [[c, money(v), f"{v / wpcc_tot['total']:.2%}".replace(".", ",")]
            for c, v in zip(CATS, wpcc_tot["cats"])]
    rows.append(["ΣΥΝΟΛΟ", money(wpcc_tot["total"]), "100,00%"])
    D.table(["Κατηγορία δαπάνης", "Ποσό (EUR)", "Ποσοστό"], rows,
            widths=[8.5, 4.5, 3.5])

    # -------------------------------------------------------- work packages
    D.h("4. Εγκεκριμένος προϋπολογισμός κατά πακέτο εργασίας")
    rows = []
    for w in wps:
        off = "" if (w["start"], w["end"]) == (hdr["start"], hdr["end"]) \
            else "διαφορετικές από την πράξη"
        rows.append([f"ΠΕ{w['wp']}", w["title"], w["start"], w["end"],
                     money(w["budget"]), off])
    rows.append(["", "ΣΥΝΟΛΟ", "", "", money(wp_grand), ""])
    D.table(["Α/Α", "Τίτλος πακέτου εργασίας", "Έναρξη", "Λήξη",
             "Ποσό (EUR)", "Ημερομηνίες"], rows,
            widths=[1.2, 6.2, 2.0, 2.0, 2.4, 3.0], right_from=4)

    D.h("5. Πακέτο εργασίας ανά κατηγορία δαπάνης")
    rows = []
    for r in wpcc:
        rows.append([f"ΠΕ{r['wp']}"] + [money(v) for v in r["cats"]]
                    + [money(r["total"])])
    rows.append(["ΣΥΝΟΛΑ"] + [money(v) for v in wpcc_tot["cats"]]
                + [money(wpcc_tot["total"])])
    D.table(["ΠΕ"] + CATS + ["ΣΥΝΟΛΟ"], rows,
            widths=[1.1] + [2.1] * 6 + [2.2])

    # ------------------------- the analytical table, one per partner
    D.h("6. Αναλυτικός προϋπολογισμός κατά δικαιούχο, παραδοτέο και "
        "κατηγορία δαπάνης")
    D.p("Ο πίνακας που ακολουθεί είναι η αναλυτικότερη μορφή στην οποία ο "
        "εγκεκριμένος προϋπολογισμός υπάρχει στο ΟΠΣ. Κάθε γραμμή είναι ένα "
        "παραδοτέο ενός εταίρου, αναλυμένο στις έξι κατηγορίες δαπάνης.",
        size=9, colour=GREY)
    for p in parts:
        f = fund_rows[p["no"] - 1] if p["no"] <= len(fund_rows) else {}
        D.h(f"6.{p['no']}  {p['name']}  ({f.get('country', '')})", size=11,
            space=8)
        rows = []
        for r in p["rows"]:
            rows.append([r["deliv"], f"ΠΕ{r['wp']}", r["title"]]
                        + [money(v) for v in r["cats"]]
                        + [money(r["total"])])
        rows.append(["ΣΥΝΟΛΟ", "", ""] + [""] * 6
                    + [money(p["printed_total"] or 0)])
        D.table(["Παραδοτέο", "ΠΕ", "Τίτλος"] + CATS + ["ΣΥΝΟΛΟ"], rows,
                widths=[1.5, 1.0, 4.2] + [1.6] * 6 + [1.9], right_from=3)

    # ------------------------------------------------------------ indicators
    if inds:
        D.h("7. Δείκτες")
        rows = [[d["code"], d["title"],
                 "εκροής" if d["kind"] == "output" else "αποτελέσματος",
                 d["unit"], str(d["target"])] for d in inds]
        D.table(["Κωδικός", "Τίτλος", "Τύπος", "Μονάδα", "Στόχος"], rows,
                widths=[1.8, 7.2, 2.6, 2.0, 1.6], right_from=4)

    # ----------------------------------------------------- reporting calendar
    def gd(s):
        a = s.split("/")
        return dt.date(int(a[2]), int(a[1]), int(a[0]))

    sched, sched_notes = schedule(mis, gd(hdr["start"]),
                                  gd(hdr["end"]))
    if sched:
        D.h("8. Ημερολόγιο περιόδων αναφοράς")
        D.table(["Έκθεση", "Από", "Έως", "Διάρκεια", "Παρατήρηση"], sched,
                widths=[2.2, 2.4, 2.4, 2.4, 6.0], right_from=1)

    # ------------------------------------------- reviewer block, HIS rule
    D.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    D.h("ΣΗΜΕΙΩΣΕΙΣ ΓΙΑ ΕΣΩΤΕΡΙΚΗ ΧΡΗΣΗ, ΝΑ ΔΙΑΓΡΑΦΟΥΝ ΠΡΙΝ ΤΗΝ ΠΡΟΩΘΗΣΗ",
        size=12, colour=RED)

    notes = [
        ("Αυτό είναι ο ΕΓΚΕΚΡΙΜΕΝΟΣ ΠΡΟΫΠΟΛΟΓΙΣΜΟΣ, όχι δαπάνες",
         "Κάθε ποσό στο παράρτημα προέρχεται από το Τεχνικό Δελτίο και "
         "δηλώνει τι επιτρέπεται να δαπανηθεί. Δεν αποτελεί απορρόφηση ούτε "
         "πιστοποιημένη δαπάνη. Αν τοποθετηθεί κάτω από επικεφαλίδα που λέει "
         "δαπάνες, ο αναγνώστης θα συμπεράνει ότι τα ποσά έχουν εκταμιευθεί."),
    ]

    # THE COINCIDENCE THAT WILL BE MISREAD IF NOBODY NAMES IT
    #
    # On an 80/20 programme, 20% of the total eligible cost is numerically
    # IDENTICAL to the national contribution. Both are 177.557,25 here. Two
    # completely different quantities, one number, and nothing on the page
    # distinguishes them - so a reader who meets 177.557,25 in a paragraph
    # cannot tell which one is meant, and neither can the person writing it
    # three weeks from now.
    twenty = round(fund_tot["total"] * 0.20, 2)
    if abs(twenty - fund_tot["nat"]) < 0.02:
        notes.append((
            "Προσοχή: δύο διαφορετικά μεγέθη με το ίδιο ποσό",
            f"Το 20% της συνολικής επιλέξιμης δαπάνης είναι "
            f"{money(twenty)}. Η εθνική συμμετοχή είναι επίσης "
            f"{money(fund_tot['nat'])}, επειδή η συγχρηματοδότηση είναι "
            f"80/20. Πρόκειται για δύο εντελώς διαφορετικά μεγέθη με "
            f"ταυτόσημη αριθμητική τιμή. Όπου αναφέρεται το ποσό, να "
            f"προσδιορίζεται ρητά ποιο από τα δύο εννοείται."))

    off = [w for w in wps
           if (w["start"], w["end"]) != (hdr["start"], hdr["end"])]
    if off:
        notes.append((
            "Πακέτα εργασίας με ημερομηνίες διαφορετικές από την πράξη",
            "; ".join(f"ΠΕ{w['wp']} ({w['start']} έως {w['end']})"
                      for w in off)
            + f". Η πράξη διαρκεί {hdr['start']} έως {hdr['end']}. Στην "
            f"τροποποίηση που μετέθεσε την περίοδο υλοποίησης, τα υπόλοιπα "
            f"πακέτα μετατοπίστηκαν και αυτό όχι. Να διευκρινιστεί αν είναι "
            f"σκόπιμο ή παράλειψη, πριν το ερώτημα τεθεί από τη Διαχειριστική "
            f"Αρχή."))

    notes.extend(sched_notes)
    notes.append((
        "Τι ΔΕΝ περιέχει αυτό το παράρτημα",
        "Δεν περιέχει φυσική πρόοδο, συμβασιοποίηση, ούτε πιστοποιημένη "
        "δαπάνη. Οι σχετικοί πίνακες και τα γραφήματα παράγονται από το Πλάνο "
        "Συμβάσεων με το build_charts.py και τα πραγματικά στοιχεία "
        "αφήγησης με το build_narrative.py."))

    for title, body in notes:
        D.h(title, size=10.5, colour=RED, space=8)
        D.p(body, size=9.5)

    out = HERE / (f"{hdr['acronym']}_Budget_Annex_v{hdr['version']}_"
                  f"{dt.date.today():%Y-%m-%d}.docx")
    D.d.save(out)
    # LAST LINE, ALWAYS. Reads the finished file back and deletes it if the
    # other programme appears anywhere in it, paragraphs or table cells.
    guard(out, forbidden)
    print(f"wrote {out}")
    print(f"  {len(parts)} partners, "
          f"{sum(len(p['rows']) for p in parts)} deliverable rows, "
          f"total {money(fund_tot['total'])}")
    print(f"  {len(notes)} reviewer note(s) behind the page break")
    return out


if __name__ == "__main__":
    main(pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else None)
